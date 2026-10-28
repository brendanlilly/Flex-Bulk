"""Bulk order (v1.2): BOM upload, multi-item search and quick entry -> personalized, stock-aware matches.

Per line:
  1. parse      resolve.parse_table: a header row anywhere in the first rows maps qty / part / description /
                manufacturer / UOM columns; without one, lines are read like a counter paste. Excel and
                text PDFs are turned into text first (read_upload).
  2. identify   resolve.Resolver: catalog exact -> real Webb item -> cross-reference. No guessing yet.
  3. recommend  no exact hit: search the description (or the line) and re-rank the relevant candidates by
                this account's own context, with the reasons shown on each:
                  - bought before (times, last date, the job it went to)
                  - last invoiced price and invoice #
                  - preferred brand for that product type (share of the account's lines)
                  - on hand at the pickup branch for the quantity asked
  4. fulfil     quantity vs on hand at the pickup branch: pickup today, part shelf + transfer, all transfer
                (which branches, how many, ETA) or CD / special order. When the pick is short, same-spec
                substitutes that are on the shelf (previously bought and usual brand first); for real Webb
                items, cross-reference equivalents in stock.
  5. confirm    exact matches arrive confirmed; recommendations, cross-reference picks and substitutes need
                the user's click, and a review screen confirms the whole order before the cart (bulk.js).

Guests get matching and branch stock, no personal ranking. Technicians get their own history, no prices
and no invoices, the same limits as everywhere else.
"""
from __future__ import annotations

import base64
import copy
import io
import json
import math
import re
import time
from collections import Counter

import numpy as np

from webbsearch.index.fixture_rules import implausible
from webbsearch.retrieval.analyzer import compact_id
from webbsearch.retrieval.availability import fulfil
from webbsearch.retrieval.resolve import MAX_LINES, Resolver, parse_table, part_tokens
from webbsearch.merch.bulkprefs import brand_cross, brand_rank, item_cross, line_key, other_brand, replace_brand
from webbsearch.merch.rules import bulk_weights
from webbsearch.merch.store import DEFAULTS as _MERCH_DEFAULTS
from webbsearch.retrieval import linefit, trade_terms
from webbsearch.retrieval.shorthand import expand, ocr_repair

# pipe schedules, pressure classes and ASTM specs read like part numbers but describe the item
SPEC_TOKEN = re.compile(r"^(?:SCH|SCHED|CLASS|CL|TYPE)[-#]?\d+$|^(?:A53|A105|A106|A234|A312|A197|A47|A126|A48|B88|B16|B62|B584)$", re.I)
DEFAULT_BULK = dict(_MERCH_DEFAULTS["bulk"])
MAX_UPLOAD = 15_000_000     # phone photos and scanned PDFs run 3-12 MB
FAMILY_LABEL = lambda f: (f or "").replace("_", " ")


# ------------------------------------------------------------------ files
def read_upload(filename: str, data_b64: str, sheet: str | None = None) -> dict:
    """An uploaded BOM -> text parse_table can read. CSV/TSV/TXT as is; Excel's chosen (or first non-empty)
    sheet as tab-separated rows; a text PDF's lines. Scanned (image) PDFs have no text to read."""
    raw = base64.b64decode(data_b64.split(",", 1)[-1] if data_b64.startswith("data:") else data_b64)
    if len(raw) > MAX_UPLOAD:
        raise ValueError(f"file is {len(raw) / 1e6:.1f} MB; the limit is 15 MB")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        names = wb.sheetnames
        pick = sheet if sheet in names else None
        rows_by = {}
        for sn in ([pick] if pick else names):
            ws = wb[sn]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                         for v in row]
                while cells and not cells[-1].strip():
                    cells.pop()
                if any(c.strip() for c in cells):
                    rows.append("\t".join(c.replace("\t", " ").replace("\n", " ") for c in cells))
                if len(rows) > MAX_LINES + 40:
                    break
            rows_by[sn] = rows
            if rows and not pick:
                pick = sn
                break
        wb.close()
        return {"kind": "xlsx", "sheets": names, "sheet": pick, "text": "\n".join(rows_by.get(pick) or [])}
    if name.endswith(".pdf") or raw[:4] == b"%PDF":
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(raw))
        pages = [(pg.extract_text() or "") for pg in r.pages[:40]]
        text = "\n".join(pages)
        return {"kind": "pdf", "pages": len(r.pages), "text": text,
                "note": "" if text.strip() else "No text in this PDF (a scan?). Export the BOM as Excel or CSV instead."}
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return {"kind": "text", "text": raw.decode(enc)}
        except UnicodeDecodeError:
            continue
    return {"kind": "text", "text": raw.decode("utf-8", "replace")}


# ------------------------------------------------------------------ matching
class Bom:
    def __init__(self, engine):
        self.e = engine
        self.res = Resolver(engine)
        self._inv: dict[str, dict] = {}
        self._jobs_cache: dict[str, dict] = {}
        self._cache: dict[str, tuple] = {}

    def _brands(self) -> dict:
        """One-word catalog brands, for fixing brand names a scan misread ("Wibco" -> NIBCO)."""
        if getattr(self, "_brand_map", None) is None:
            names = set(self.e.catalog.brand.tolist())
            self._brand_map = {b.lower(): b for b in names if b and " " not in b and b.isalpha()}
        return self._brand_map

    # -------------------------------------------------------------- Flex Bulk: brand names, shorthand, near part #s
    def _brand_names(self):
        try:
            from flexbulk.brands import get
            return get(self.e)
        except Exception:  # noqa: BLE001
            return None

    def _short(self, text: str) -> str:
        try:
            from flexbulk.shorthand import get
            return get().expand(text)
        except Exception:  # noqa: BLE001
            return text

    def _brand_count(self) -> Counter:
        if getattr(self, "_bc", None) is None:
            self._bc = Counter(map(str, self.e.catalog.brand))
        return self._bc

    def near_parts(self, keys: list[str], bset) -> list[tuple[int, str]]:
        """Items of the line's brand whose part # is the line's with something added or cut ("ENS-40" ->
        ENS-40 200; "EX30B" -> EX30): (item, the part # it has). Only with a brand to check against."""
        if not bset:
            return []
        cat, out, seen = self.e.catalog, [], set()
        for k in keys:
            if len(k) < 4:
                continue
            cands = list(cat.id_prefix(k, 30)) if len(k) >= 5 else []
            for cut in (1, 2, 3):
                if len(k) - cut >= 4:
                    cands += list(cat.id_lookup(k[:-cut]))
            for j in cands:
                if j not in seen and str(cat.brand[j]) in bset:
                    seen.add(j)
                    out.append((int(j), str(cat.mfr_part[j])))
        return out[:3]

    # -------------------------------------------------------------- account context
    def _invoiced(self, p) -> dict:
        """sku -> last invoiced line for this account (price, invoice #, date, job). Needs invoice access."""
        if not (p.is_authenticated and p.can("invoices")):
            return {}
        c = self._inv.get(p.account_id)
        if c is not None:
            return c
        acc = self.e.accounts
        ids = acc.inv_by_acct.get(p.account_id)
        out: dict[str, dict] = {}
        if ids is not None and len(ids):
            inv = acc.invoices.loc[ids]
            by_order = {o: (i, d, s) for o, i, d, s in zip(inv["order_id"], inv["invoice_id"], inv["invoice_date"], inv["status"])}
            lo = acc._lo[acc._lo["account_id"] == p.account_id]
            lo = lo[lo["order_id"].isin(by_order.keys())]
            jobs = acc.orders["job_name"]
            for sku, oid, price in zip(lo["sku"], lo["order_id"], lo["unit_price"]):   # sorted by order date
                i_, d_, s_ = by_order[oid]
                out[sku] = {"price": round(float(price), 2), "invoice_id": str(i_), "date": str(d_.date()),
                            "status": s_, "job": str(jobs.get(oid) or "")}
        if len(self._inv) > 200:
            self._inv.clear()
        self._inv[p.account_id] = out
        return out

    def _jobs(self, p) -> dict:
        """sku -> the job it last went to, for this account (a technician's own orders only)."""
        if not p.is_authenticated:
            return {}
        key = p.account_id + ("|" + p.user_id if p.role == "tech" else "")
        c = self._jobs_cache.get(key)
        if c is None:
            acc = self.e.accounts
            ids = acc._order_ids(p)
            lo = acc._lo[acc._lo["order_id"].isin(set(ids.tolist()))]
            jobs = acc.orders["job_name"]
            c = {sku: str(jobs.get(oid) or "") for sku, oid in zip(lo["sku"], lo["order_id"])}   # last wins (date order)
            if len(self._jobs_cache) > 200:
                self._jobs_cache.clear()
            self._jobs_cache[key] = c
        return c

    def context(self, p, branch: str | None) -> dict:
        e = self.e
        branch = branch or (p.home_branch if p.is_authenticated else None)
        out = {"signed_in": p.is_authenticated, "branch": e.availability.branch_info(branch),
               "role": p.role, "sees_prices": not p.is_authenticated or p.sees_prices}
        if not p.is_authenticated:
            return out
        out["account"] = p.account_name
        ap = e.features.acct_prod.get(p.account_id, {})
        out["skus_bought"] = len(ap)
        brands = []
        for fam, share in e.features.top_families(p.account_id, 8):
            pb = e.features.preferred_brand(p.account_id, fam)
            if pb and pb[1] >= 0.35:
                brands.append({"family": FAMILY_LABEL(fam), "brand": pb[0], "share": round(pb[1], 2)})
        out["preferred_brands"] = brands[:5]
        out["invoiced_skus"] = len(self._invoiced(p))
        return out

    # -------------------------------------------------------------- cards
    def _card(self, j: int, p, branch, need: int, inv: dict, pref: dict, uom: str = "", ctx: dict | None = None) -> dict:
        e = self.e
        cat = e.catalog
        m = e.results._mini(int(j), p)
        sku = m["sku"]
        m.update(kind="catalog", mfr_part=cat.mfr_part[j], list_price=float(cat.list_price[j]),
                 stock_tier=cat.stock_tier[j])
        if p.is_authenticated and not p.sees_prices:
            m["price"] = m["list_price"] = None
        # "200 FT of 1/2 PEX" against a 100 ft coil is 2 coils: order and check stock in the item's own unit
        L = re.search(r"(\d+(?:\.\d+)?)\s*ft\b", cat.title[j], re.I) if uom in ("FT", "LF", "FEET") else None
        if L and float(L.group(1)) > 0:
            units = max(1, math.ceil(need / float(L.group(1))))
            m["units"] = units
            m["conv"] = f"{need:,} ft = {units} × {L.group(1)} ft"
            need = units
        m["fulfil"] = fulfil(need, e.availability.all_qty(int(j)), e.availability, branch,
                             "special" if cat.stock_tier[j] == "special" else "dc")
        why, bonus, parts = [], 0.0, {}
        ms = (ctx or {}).get("ms") or (e.merch.get() if getattr(e, "merch", None) else None)
        bw = bulk_weights(ms) if ms else DEFAULT_BULK
        h = e.accounts.history_badge(p, int(j))
        if h:
            job = self._jobs(p).get(sku, "")
            why.append({"k": "hist", "t": f"Bought {h['times']}× · last {h['last_date']}" + (f" · {job}" if job else "")})
            bonus += bw["bought_before"] + bw["bought_times"] * math.log1p(h["times"])
            parts["bought before"] = bw["bought_before"] + bw["bought_times"] * math.log1p(h["times"])
            m["history"] = h
        iv = inv.get(sku)
        if iv:
            why.append({"k": "inv", "t": (f"Last invoiced {'$' + format(iv['price'], ',.2f') if p.sees_prices else ''}"
                                          f" · INV {iv['invoice_id']} ({iv['date']})").replace("invoiced  ·", "invoiced ·")})
            bonus += bw["invoiced"]
        fam = cat.family[j]
        shares = pref.get(fam)
        if shares is None and p.is_authenticated:
            shares = pref[fam] = self.e.features.acct_brand.get((p.account_id, fam), {})
        if shares:
            s = shares.get(cat.brand[j], 0.0)
            top = max(shares, key=shares.get)
            if top == cat.brand[j] and s >= 0.35:
                why.append({"k": "brand", "t": f"Your usual {FAMILY_LABEL(fam)} brand ({s:.0%} of your lines)"})
                bonus += bw["preferred_brand"] * s
            elif s >= 0.15:
                bonus += 0.4 * bw["preferred_brand"] * s
        f = m["fulfil"]
        if f["status"] == "pickup":
            bonus += bw["on_shelf"]
            parts["on the shelf"] = bw["on_shelf"]
        elif f["at_branch"]:
            bonus += 0.25 * bw["on_shelf"]
            parts["some on the shelf"] = 0.25 * bw["on_shelf"]
        if ctx:
            ms = ms or {}
            # Webb's brand order for the category, when the line names no brand (a branch's own order first)
            if not ctx.get("named_brand") and bw.get("webb_brand"):
                rk = brand_rank(ms, branch, str(cat.family[j]), str(cat.cat3[j]), str(cat.brand[j]))
                if rk:
                    r_, n_ = rk
                    bonus += bw["webb_brand"] * (n_ - r_) / n_
                    parts[f"Webb brand #{r_ + 1}"] = bw["webb_brand"] * (n_ - r_) / n_
                    why.append({"k": "webb", "t": f"Webb's #{r_ + 1} brand for {str(cat.cat3[j]).lower()}"})
            lr = ctx.get("learned")
            if lr is not None:
                b_, w_ = lr.bonus(bw, ctx.get("off") or set(), ctx["key"], sku, str(cat.brand[j]), str(cat.family[j]),
                                  ctx.get("req_brand") or "", ctx.get("acct") or p.account_id)
                if b_:
                    bonus += b_
                    parts["learned"] = b_
                    why.append({"k": "learned", "t": w_})
        m["why"] = why
        m["_bonus"] = bonus
        m["_parts"] = {k: round(v, 2) for k, v in parts.items()}
        return m

    def _webb_card(self, c: dict, need: int, branch) -> dict:
        """A cross-reference or real Webb item card (no demo price) with its fulfilment."""
        e = self.e
        out = {"kind": "webb", "sku": c["sku"], "title": c["title"], "brand": c.get("brand", ""),
               "mfr_part": c.get("mpn") or c.get("mfr_part") or "", "webb_item": c["webb_item"], "url": c.get("url"),
               "price": c.get("price") if c.get("price") is not None else c.get("list_price"), "uom": "EA",
               "family": "boiler_parts", "attrs": {}, "noun": c.get("noun", ""), "why": []}
        for k in ("band", "wmcr", "how"):
            if c.get(k) is not None:
                out[k] = c[k]
        part = e.xref.webb_part(c["webb_item"]) if e.xref.ok and e.xref.stock else None
        if part is not None:
            out["cls"] = part["cls"]
            out["fulfil"] = fulfil(need, e.xref.stock.all_qty(part["webb_item"], part["cls"]), e.availability, branch)
            out["fulfil"]["demo"] = True
        return out

    # -------------------------------------------------------------- recommendations
    def recommend(self, query: str, p, branch, need: int, k: int = 6, uom: str = "", ctx: dict | None = None) -> tuple[list[dict], bool, dict]:
        """Relevant candidates for a description, re-ranked by this account's history, invoices, preferred
        brand and branch stock. Returns (cards, weak, fit): candidates of the product the line names come first
        (a wye before the no-hub coupling that merely shares words), then the right size, then score. When the
        first search has none of that product in that size, a short focused search runs too. `fit` says whether
        the catalog had the product at all and what the top pick misses (size, material, joint)."""
        e = self.e
        sp = linefit.spec(query, uom)
        ms = (ctx or {}).get("ms")             # the draft's settings when the dashboard tries lines, else live
        tr = (ctx or {}).get("trace")
        res = e.results.results(query, p, branch=branch, llm="off", settings=ms)
        nt = res.get("notice") or {}
        if tr is not None:
            tr.append({"step": "search", "query": query, "results": res.get("total"), "notice": nt.get("kind"),
                       "spec": {"product": sp.head, "size": sp.size, "material": sp.material, "joint": sp.conn,
                                "angle": sp.angle, "reducing": sp.reducing}})
        bset = (ctx or {}).get("brands") or frozenset()
        if nt.get("kind") == "not_carried" and not bset:
            # the catalog doesn't carry the product (a hydrant, C900 pipe): no recommendation, no focused re-search
            return [], True, {"head": sp.head, "not_carried": nt.get("what") or query}
        items = [] if nt.get("kind") == "not_carried" else (res.get("items") or [])[:24]
        if bset and not any(str(it.get("brand")) in bset for it in items):
            # the line names a manufacturer the list has: look among its items before saying it isn't carried
            cb = max(bset, key=lambda b: self._brand_count().get(b, 0))
            bwords = set(re.findall(r"[a-z0-9]+", " ".join(bset).lower() + " " + (ctx or {}).get("line_brand", "").lower()))
            content = [w for w in re.findall(r"[A-Za-z][A-Za-z-]{2,}", query) if w.lower() not in bwords]
            tries = [query if cb.lower() in query.lower() else f"{cb} {query}"]
            if sp.head:
                tries.append(f"{cb} {sp.head}")
            if content:
                tries.append(f"{cb} {' '.join(content[-3:])}")
            seen_ = {x["sku"] for x in items}
            for bq in dict.fromkeys(tries):
                more = [it for it in (e.results.results(bq, p, branch=branch, llm="off", settings=ms).get("items") or [])[:24]
                        if str(it.get("brand")) in bset and it["sku"] not in seen_]
                if tr is not None:
                    tr.append({"step": "brand search", "query": bq, "found": len(more)})
                if more:
                    items = items + more
                    break
            if nt.get("kind") == "not_carried" and not items:
                return [], True, {"head": sp.head, "not_carried": nt.get("what") or query}
        weak = bool(res.get("notice") and res["notice"].get("kind") in ("fallback", "partial", "phrase_miss"))
        fits = {it["sku"]: linefit.fit(sp, it) for it in items}
        # none of the right product, size and spec on the first page: search for just that; then without the
        # size (a 1/2" trap primer the catalog only has in 2", say, is still the nearest)
        for with_size in (True, False):
            if not sp.head or any(f["good"] for f in fits.values()) or (not with_size and any(f["noun"] for f in fits.values())):
                break
            fq = linefit.focus_query(sp, with_size)
            if fq and fq.lower() != query.lower():
                if tr is not None:
                    tr.append({"step": "focused search", "query": fq, "why": "none of the first page is the line's product"
                               + (" in its size" if with_size else "")})
                seen = {it["sku"] for it in items}
                more = [it for it in (e.results.results(fq, p, branch=branch, llm="off", settings=ms).get("items") or [])[:24]
                        if it["sku"] not in seen]
                items += more
                fits.update({it["sku"]: linefit.fit(sp, it) for it in more})
        inv, pref = self._invoiced(p), {}
        # a recommendation claims the item is right: skip catalog records a tradesperson would call impossible
        ok = [it for it in items if not implausible(it["title"])]
        items = ok or items
        cards = []
        for pos, it in enumerate(items):
            j = e.catalog.sku_to_idx.get(it["sku"])
            if j is None:
                continue
            c = self._card(j, p, branch, need, inv, pref, uom, ctx)
            f = fits[it["sku"]]
            c["exact"] = bool(it.get("exact"))
            noun, size, good = (True, True, True) if c["exact"] else (f["noun"], f["size"], f["good"])
            # the product the line names, in its size, as specified (street, 45, reducing...) first; then the
            # account's history, invoices, brand and stock, and search position
            bonus_ = c.pop("_bonus")
            brand_ok = not bset or str(c.get("brand")) in bset or bool(c.get("exact"))
            c["_key"] = (not noun, not brand_ok, not size, not good, -(10 * (1 - pos / max(len(items), 1)) + bonus_ + f["score"]))
            c["_fit"] = {"good": True, "misses": []} if c["exact"] else f
            if tr is not None:
                c["_trace"] = {"sku": c["sku"], "title": c["title"], "brand": c.get("brand"), "search_pos": pos + 1,
                               "product": noun, "brand_ok": brand_ok, "size": size, "spec_ok": good, "misses": ", ".join(f.get("misses") or []),
                               "search": round(10 * (1 - pos / max(len(items), 1)), 2), "fit": round(f["score"], 2),
                               "bonus": round(bonus_, 2), "bonus_parts": c.get("_parts") or {},
                               "total": round(-c["_key"][4], 2)}
            cards.append(c)
        cards.sort(key=lambda c: c["_key"])
        if tr is not None:
            tr.append({"step": "ranked", "rule": "the line's product first, then the line's brand, then its size, then nothing "
                       "missing, then search position + fit + bonus", "candidates": [dict(c.pop("_trace"), rank=k + 1) for k, c in enumerate(cards[:15])]})
            for c in cards[15:]:
                c.pop("_trace", None)
        top = cards[0]["_fit"] if cards else {"good": False, "misses": []}
        info = {"head": sp.head, "head_label": linefit.head_label(sp), "size": sp.size, "what": linefit.describe(sp),
                "missing": bool(sp.head and cards and cards[0]["_key"][0]), "misses": top["misses"] + top.get("soft", []),
                "good": bool(top["good"])}
        if sp.head and top["good"]:
            weak = False            # the line's product, size and spec: unmatched filler words don't make it weak
        for c in cards:
            c.pop("_key", None)
            c.pop("_fit", None)
            c.pop("_parts", None)
        return cards[:k], weak, info

    def substitutes(self, j: int, p, branch, need: int, k: int = 3, uom: str = "", ctx: dict | None = None) -> list[dict]:
        """Same-spec items (every attribute but the series) that are on the shelf for the quantity asked,
        previously bought and usual brand first."""
        e = self.e
        cat = e.catalog
        a = cat.attrs(int(j))
        key = {x: v for x, v in a.items() if x not in ("series",)}
        if not key:
            return []
        pq = e.parser.parse(" ".join(w for w in cat.title[j].split()[1:]))
        c2 = e.retriever.retrieve(pq, filters={"family": [cat.family[j]]}, lanes=("lexical",), limit=150)
        same = [int(x) for x in c2.idx.tolist() if int(x) != int(j) and all(cat.attrs(int(x)).get(kk) == v for kk, v in key.items())]
        if not same:
            return []
        av = e.availability.summary(np.array(same, np.int64), branch)
        ok = [s for s, at in zip(same, av["at_branch"].tolist()) if at > 0]
        inv, pref = self._invoiced(p), {}
        cards = [c for c in (self._card(s, p, branch, need, inv, pref, uom, ctx) for s in ok[:20]) if c["fulfil"]["status"] == "pickup"]
        price0 = float(cat.list_price[j])
        cards.sort(key=lambda c: (-c["_bonus"], abs((c["list_price"] or price0) - price0)))
        for c in cards:
            c.pop("_bonus", None)
            c.pop("_parts", None)
            c["how"] = f"Same spec, {c['fulfil']['at_branch']} on hand at {c['fulfil']['branch']}"
        return cards[:k]

    # -------------------------------------------------------------- one line
    def match(self, ln: dict, p, branch, k: int = 6, settings: dict | None = None, acct: str | None = None,
              trace: bool = False) -> dict:
        """`settings`: merchandising settings to match with instead of the live ones (the dashboard's draft test).
        `acct` (Flex Bulk): the customer the order is for, so their own past choices count for more."""
        from webbsearch.retrieval.analyzer import clean_dashes      # en / em dashes from Word and PDFs read as hyphens
        ln = {**ln, **{k: clean_dashes(ln[k]) for k in ("text", "desc", "part", "brand") if isinstance(ln.get(k), str)}}
        e = self.e
        branch = branch or (p.home_branch if p.is_authenticated else None)
        need = max(1, int(ln.get("qty") or 1))
        uom = (ln.get("uom") or "").upper().rstrip(".")
        # merchandising crosses and learned choices (bulkprefs.py)
        ms = settings or (e.merch.get() if getattr(e, "merch", None) else {})
        line_text = (ln.get("text") or "") + " " + (ln.get("desc") or "")
        key = line_key(expand(ln.get("desc") or ln.get("text") or "", self._brands()))
        bc = brand_cross(ms, branch, line_text, ln.get("brand") or "")
        parsed_brands = e.parser.parse(line_text[:300]).brands
        uncarried = "" if bc or parsed_brands else other_brand(line_text)
        req_brand = bc["from_brand"] if bc else ((ln.get("brand") or "").strip() or (parsed_brands[0] if parsed_brands else "") or uncarried)
        # Flex Bulk: the list's brands the line's manufacturer means (brand names table, flexbulk/brands.py)
        bn = self._brand_names()
        line_brand = (ln.get("brand") or "").strip() or (parsed_brands[0] if parsed_brands else "")
        bset = bn.resolve(line_brand) if bn is not None and line_brand and not bc else frozenset()
        ctx = {"key": key, "req_brand": req_brand, "named_brand": (bool(req_brand) and not bc and not uncarried) or bool(bset),
               "brands": bset, "line_brand": line_brand,
               "learned": getattr(e, "learned", None), "off": set(ms.get("learned_off") or []), "ms": ms, "acct": acct,
               "trace": [] if trace else None}
        tr = ctx["trace"]
        if tr is not None:
            tr.append({"step": "read", "text": ln.get("text"), "desc": ln.get("desc"), "brand": ln.get("brand"),
                       "part": ln.get("part"), "qty": need, "uom": uom, "from_ocr": bool(ln.get("from_ocr")),
                       "line_key": key, "brand_named": req_brand or None, "brand_on_list": sorted(bset)[:6] or None})
        pkeys = {re.sub(r"[^a-z0-9]", "", t.lower()) for t in part_tokens(ln.get("part") or "") + part_tokens(ln.get("text") or "")}
        ic = item_cross(ms, branch, key, pkeys)
        if ic and ic["to_sku"] in e.catalog.sku_to_idx:
            r0 = {"status": "xref", "via": "Webb cross (set by merchandising)", "item": {"kind": "catalog", "sku": ic["to_sku"]},
                  "alternates": [], "note": f"Webb cross: {ic['label']}" + (f" · {ic['note']}" if ic.get("note") else "")}
        else:
            r0 = self.res.one(ln, p, branch, best=False)
        mismatch = []                       # Flex Bulk: the line's part # is another manufacturer's item
        if r0["status"] == "exact" and bset and r0.get("item"):
            items = [r0["item"]] + list(r0.get("alternates") or [])
            same = [it for it in items if str(it.get("brand")) in bset]
            if same:
                if same[0] is not items[0]:
                    r0 = {**r0, "item": same[0], "alternates": [it for it in items if it is not same[0]],
                          "note": f"{len(items)} items share this number; the {line_brand} one first"}
            else:
                mismatch = [it for it in items if it.get("kind") == "catalog"]
                r0 = {"status": "none", "via": None, "item": None, "alternates": [], "note": ""}
        if tr is not None:
            if mismatch:
                tr.append({"step": "brand check", "found": f"part # is {mismatch[0].get('brand')}'s, the line names {line_brand}"})
            tr.append({"step": "item cross" if ic else "brand cross" if bc else "crosses", "found": (ic or bc or {}).get("label")
                       or (f"{bc['from_brand']} -> {bc['to_brand']}" if bc else None)})
            tr.append({"step": "exact lookup", "result": r0["status"], "via": r0.get("via"), "note": r0.get("note") or None,
                       "item": (r0.get("item") or {}).get("sku")})
        out = {f: ln.get(f) for f in ("n", "raw", "text", "desc", "brand", "part", "qty", "uom", "qty_how",
                                      "src", "src_label", "crop", "conf", "from_ocr", "ocr_raw", "page")}
        out.update(status=r0["status"], via=r0["via"], note=r0.get("note") or "", options=[], sel=0,
                   confirmed=False, needs_confirm=False, substitutes=[], closest=[], xref_part=r0.get("xref_part"),
                   key=key, req_brand=req_brand, uncarried=bool(uncarried) or bool(bc))
        if ic and (r0.get("via") or "").startswith("Webb cross"):
            out["cross"] = {"id": ic["id"], "label": ic["label"]}
        inv, pref = self._invoiced(p), {}
        if r0["status"] == "skip":
            return out
        if r0["status"] in ("exact", "webb", "xref") and r0["item"]:
            items = [r0["item"]] + (r0.get("alternates") or [])
            for it in items[:k]:
                if it.get("kind") == "catalog":
                    j = e.catalog.sku_to_idx.get(it["sku"])
                    if j is not None:
                        out["options"].append(self._card(j, p, branch, need, inv, pref, uom, ctx))
                        out["options"][-1].pop("_bonus", None)
                else:
                    out["options"].append(self._webb_card(it, need, branch))
            if r0["status"] == "xref":
                out["options"][0].setdefault("band", r0.get("band"))
                out["options"][0].setdefault("wmcr", r0.get("wmcr"))
            # a real Webb item found by brand + description words is a recommendation, not an exact match
            out["confirmed"] = r0["status"] in ("exact", "webb") and "description" not in (r0.get("via") or "")
            out["needs_confirm"] = not out["confirmed"]
            if r0.get("swap"):      # the cross-reference pick is out at the branch; an in-stock one exists
                out["substitutes"].append(self._webb_card(r0["swap"], need, branch))
        else:
            # no exact hit: recommend from the description (BOM column) or the line itself
            out["note"] = ""        # the resolver's "looks like a part number" is re-said below only if it still is one
            pts = [t for t in part_tokens(ln.get("text") or "") if " " not in t]
            # trade shorthand that looks like a part number (HD2000, 3M) is a description once expanded
            exp_line = expand(ln.get("text") or "", self._brands()).lower()
            pts = [t for t in pts if t.lower() in exp_line and not SPEC_TOKEN.match(t) and not trade_terms.is_term(t)]
            desc = (ln.get("desc") or "").strip()
            words = [w for w in (ln.get("text") or "").split() if w not in pts]
            cbrand = max(bset, key=lambda b: self._brand_count().get(b, 0)) if bset else ""
            query = " ".join(x for x in ((cbrand or ln.get("brand") or ""), desc) if x) if desc else \
                (" ".join(words) if pts and len(words) >= 2 else ln.get("text") or "")
            if cbrand and not desc and line_brand and line_brand in query and cbrand != line_brand:
                query = query.replace(line_brand, cbrand)          # "Mueller Industries ..." -> "Mueller ..."
            query = self._short(query)                          # ERP shorthand: MXG VLV ADJ -> mixing valve adjustable
            partish = bool(pts) and not desc and len(words) < 2
            if ln.get("from_ocr") and not ln.get("edited"):
                query = ocr_repair(query)          # 14 -> 1¼ on a small-bore line, +5 -> 45, IDS -> IPS ...
            if bc:
                # a brand Webb doesn't carry: match the brand merchandising crosses it to
                query = replace_brand(query, bc["from_brand"], bc["to_brand"])
                if bc["to_brand"].lower() not in query.lower():
                    query = f"{bc['to_brand']} {query}"
                out["brand_cross"] = {"id": bc["id"], "from": bc["from_brand"], "to": bc["to_brand"]}
            searched = expand(query, self._brands())
            if searched.lower() != query.lower():
                out["searched_as"] = searched
            if tr is not None:
                tr.append({"step": "description", "query": query, "expanded": searched if searched != query else None,
                           "part_numbers_on_line": pts, "searched": not (partish or not query.strip())})
            cards, weak, fit = ([], False, {}) if partish or not query.strip() else self.recommend(searched, p, branch, need, k, uom, ctx)
            # Flex Bulk: a close part # of the line's brand, and a part # match of another brand, as options to check
            pkeys_ = [compact_id(x) for x in ([ln.get("part")] if ln.get("part") else []) + pts if x]
            near = [] if mismatch else self.near_parts(list(dict.fromkeys(pkeys_)), bset)
            extra = []
            for j, mp in near:
                c = self._card(j, p, branch, need, inv, pref, uom, ctx)
                c.pop("_bonus", None)
                c.pop("_parts", None)
                c["how"] = f"Close part #: {mp}"
                extra.append(c)
            for it in mismatch:
                j = e.catalog.sku_to_idx.get(it.get("sku"))
                if j is not None:
                    c = self._card(j, p, branch, need, inv, pref, uom, ctx)
                    c.pop("_bonus", None)
                    c.pop("_parts", None)
                    c["how"] = f"Part # match, but {c.get('brand')}'s"
                    extra.append(c)
            if extra:
                dwords = set(re.findall(r"[a-z]{3,}", searched.lower()))
                def overlap(c):
                    t = set(re.findall(r"[a-z]{3,}", (c["title"] + " " + " ".join(map(str, (c.get("attrs") or {}).values()))).lower()))
                    return len(dwords & t) / max(len(dwords), 1)
                lead = [c for c in extra if c["how"].startswith("Close") or overlap(c) >= 0.25 or not cards]
                tail = [c for c in extra if c not in lead]
                skus = {c["sku"] for c in extra}
                cards = lead + [c for c in cards if c["sku"] not in skus][:max(1, k - len(extra))] + tail
                if out.get("status") == "none" or fit.get("missing") or fit.get("not_carried"):
                    fit = {**fit, "missing": False, "not_carried": None}
                out["part_note"] = (f"Close part #: {near[0][1]} for {ln.get('part') or (pts[0] if pts else '')}" if near and cards[0] is extra[0]
                                    else f"Part # {ln.get('part') or (pts[0] if pts else '')} is {mismatch[0].get('brand')}'s, the line says {line_brand}: check it"
                                    if mismatch else "")
                pts = []                               # the part # was used: don't say "no item with part #"
                if tr is not None:
                    tr.append({"step": "part # options", "near": [mp for _, mp in near], "other_brand": [it.get("sku") for it in mismatch],
                               "first": cards[0]["sku"]})
            if fit.get("not_carried"):
                out["status"], out["via"] = "none", "Not in catalog"
                out["note"] = f"No {fit['not_carried']} in this catalog."
                out["missing_what"] = fit["not_carried"]
                out["needs_confirm"] = False
            if cards and fit.get("missing") and not (bset and any(str(c.get("brand")) in bset for c in cards)):
                # none of the candidates is the product the line names: say so, keep the nearest to pick from
                out["status"], out["via"] = "none", "Not in catalog"
                out["note"] = f"No {fit['what']} in this catalog."
                out["missing_what"] = fit["what"]
                out["closest"] = cards[:3]
                out["needs_confirm"] = False
                cards = []
            elif cards and fit.get("head") and not fit.get("good"):
                miss = [m for m in fit.get("misses", []) if m != fit.get("head")]
                out["fit_note"] = (f"Nearest match: no {' '.join(m for m in miss if m != '45')} {fit.get('head_label', '')} found" if miss and fit.get("head_label")
                                   else "Nearest match: check it against the line")
                weak = True
            if cards:
                out["options"] = cards
                exact = cards[0].get("exact") and not pts and not extra and (not bset or str(cards[0].get("brand")) in bset)
                out["status"] = "exact" if exact else "rec"
                out["via"] = "Exact" if exact else "Recommended"
                out["confirmed"] = bool(exact)
                out["needs_confirm"] = not exact
                if out.get("part_note"):
                    out["note"] = out.pop("part_note")
                    out["weak"] = True
                elif pts:
                    out["note"] = f"No item with part # {pts[0]}; recommended from “{query}”"
                    out["weak"] = True
                elif out.get("fit_note"):
                    out["note"] = out.pop("fit_note")
                    out["weak"] = True
                elif weak:
                    out["note"] = "Weak match: not every word matched"
                    out["weak"] = True
            elif not out.get("closest"):
                out["status"] = "none"
                out["note"] = out["note"] or ("Looks like a part number; no item on the list has it"
                                              if partish else "Nothing relevant found")
        if out.get("brand_cross"):
            bcx = out["brand_cross"]
            out["note"] = f"Webb crosses {bcx['from']} to {bcx['to']}" + (f" · {out['note']}" if out["note"] else "")
        if ln.get("from_ocr") and not ln.get("edited"):
            # read from handwriting or a scan: even a part-number hit may be a misread, so a person confirms it
            out["confirmed"] = False
            out["needs_confirm"] = bool(out["options"])
            out["note"] = ("Read from the note: check it against the handwriting" + (f" · {out['note']}" if out["note"] else ""))
        if tr is not None:
            for c in out["options"] + out["closest"]:
                c.pop("_parts", None)
            tr.append({"step": "result", "status": out["status"], "via": out["via"], "note": out["note"] or None,
                       "top": out["options"][0]["sku"] if out["options"] else None})
            out["trace"] = tr
        for c in out["options"]:
            c.pop("_parts", None)
        # short at the pickup branch: same-spec substitutes that are on the shelf
        if out["options"]:
            top = out["options"][0]
            f = top.get("fulfil")
            if f and f["status"] != "pickup" and top["kind"] == "catalog" and not out["substitutes"]:
                j = e.catalog.sku_to_idx.get(top["sku"])
                if j is not None:
                    out["substitutes"] = self.substitutes(j, p, branch, need, uom=uom, ctx=ctx)
            elif f and f["status"] != "pickup" and top["kind"] == "webb" and not out["substitutes"] and e.xref.ok:
                eq = e.xref.equivalents(top["webb_item"], branch, limit=6, in_stock_only=True)
                out["substitutes"] = [self._webb_card(c, need, branch) for c in eq
                                      if (c.get("stock") or {}).get("at_branch", 0) >= need][:3]
        return out

    def match_all(self, lines: list[dict], p, branch=None, k: int = 6, settings: dict | None = None,
                  acct: str | None = None, trace: bool = False) -> dict:
        """Match up to 50 lines. Identical lines (same words, quantity and unit) are matched once: takeoffs repeat
        items, and the page sends a large file in parallel chunks, so a short-lived cache spans calls too."""
        t0 = time.perf_counter()
        branch = branch or (p.home_branch if p.is_authenticated else None)
        lr = getattr(self.e, "learned", None)
        if lr is not None:
            lr.refresh()
        ver = (self.e.merch.get().get("version") if getattr(self.e, "merch", None) else 0, getattr(lr, "_mark", None))
        out, hits = [], 0
        for ln in lines[:50]:
            ck = json.dumps([p.account_id, p.user_id, acct, branch, k, ver] + [ln.get(x) for x in (
                "text", "desc", "brand", "part", "qty", "uom", "from_ocr", "edited")], default=str)
            got = None if settings is not None or trace else self._cache.get(ck)
            if got is not None and time.time() - got[0] < 900:
                m = copy.deepcopy(got[1])
                hits += 1
            else:
                m = self.match(ln, p, branch, k, settings, acct, trace)
                if settings is None and not trace:
                    if len(self._cache) > 5000:
                        self._cache.clear()
                    self._cache[ck] = (time.time(), copy.deepcopy(m))
            for x in ("n", "raw", "src", "src_label", "crop", "conf", "ocr_raw", "page", "qty_how"):
                m[x] = ln.get(x)
            out.append(m)
        counts = {s: sum(1 for x in out if x["status"] == s) for s in ("exact", "webb", "xref", "rec", "none", "skip")}
        counts["needs_confirm"] = sum(1 for x in out if x["needs_confirm"])
        counts["short"] = sum(1 for x in out if x["options"] and (x["options"][0].get("fulfil") or {}).get("status", "pickup") != "pickup")
        return {"lines": out, "counts": counts, "branch": self.e.availability.branch_info(branch),
                "personalized": p.is_authenticated, "cached": hits, "ms": round((time.perf_counter() - t0) * 1000, 1)}


def parse_request(text: str | None, filename: str | None, data_b64: str | None, sheet: str | None,
                  roles: dict | None, header_row: int | None, vocab=None, claude=None, classify: bool = False) -> dict:
    """A paste or an uploaded document -> lines to match, plus what the document said about the order.

    Files go through intake.ingest (email, Word, handwriting/scans, Excel, CSV, PDF). `roles`/`header_row` re-read
    one table with the columns the person picked. Pasted text is read as a table or a list; with `classify`
    (the BOM paste box, where people paste whole emails) non-item lines are set aside like in an email."""
    from webbsearch.retrieval.intake import free_text_lines, ingest
    if data_b64:
        raw = base64.b64decode(data_b64.split(",", 1)[-1] if data_b64.startswith("data:") else data_b64)
        if len(raw) > MAX_UPLOAD:
            raise ValueError(f"file is {len(raw) / 1e6:.1f} MB; the limit is 15 MB")
        doc = ingest(filename or "", raw, sheet=sheet, vocab=vocab, claude=claude)
        if roles:         # re-read the (single) table with the person's column choices
            src = read_upload(filename or "", data_b64, sheet) if not (filename or "").lower().endswith(".docx") else None
            table_text = src["text"] if src else "\n".join(
                __import__("webbsearch.retrieval.intake", fromlist=["read_docx"]).read_docx(raw)["tables"][:1])
            t = parse_table(table_text, MAX_LINES, roles=roles, header_row=header_row)
            s0 = doc["sources"][0] if doc["sources"] else {"id": 0, "label": filename, "kind": doc["kind"]}
            s0.update(header=t["header"], roles=t["roles"], header_row=t["header_row"],
                      lines=[{**ln, "src": s0["id"], "src_label": s0["label"], "item": True} for ln in t["lines"]])
            doc["sources"] = [s0]
        return _doc_response(doc, filename)
    t = parse_table(text or "", MAX_LINES, roles=roles, header_row=header_row)
    rows = [r for r in (text or "").replace("\r\n", "\n").split("\n") if r.strip()]
    if classify and t["header"] is None:
        s = {"id": 0, "label": "Pasted text", "kind": "text"}
        s["lines"] = free_text_lines(text or "", s, vocab)
        from webbsearch.retrieval.intake import meta_from_text
        return _doc_response({"kind": "text", "sources": [s], "meta": meta_from_text(rows), "notes": []}, None)
    out = {"kind": "text"} | t
    out["preview"] = rows[:12]
    out["filename"] = filename
    return out


def _doc_response(doc: dict, filename: str | None) -> dict:
    """intake result -> the parse response: item lines to match, the rest kept per source."""
    lines, sources = [], []
    for s in doc["sources"]:
        items = [ln for ln in s.get("lines", []) if ln.get("item", True)]
        other = [{"n": ln["n"], "raw": ln["raw"], "why": ln.get("why", ""), "text": ln["text"], "qty": ln["qty"],
                  "uom": ln.get("uom", ""), "desc": ln.get("desc", ""), "src": s["id"], "src_label": s["label"],
                  "crop": ln.get("crop"), "from_ocr": ln.get("from_ocr", False)} for ln in s.get("lines", []) if not ln.get("item", True)]
        lines += items
        sources.append({k: v for k, v in s.items() if k != "lines"} | {"n_items": len(items), "other": other})
    one = sources[0] if len(sources) == 1 else {}
    return {"kind": doc["kind"], "filename": filename, "sources": sources, "meta": doc.get("meta", {}),
            "email": doc.get("email"), "notes": doc.get("notes", []), "lines": lines[:MAX_LINES], "total": min(len(lines), MAX_LINES),
            "truncated": len(lines) > MAX_LINES, "header": one.get("header"), "roles": one.get("roles"),
            "header_row": one.get("header_row"), "sheets": one.get("sheets"), "sheet": one.get("sheet"),
            "pages": (one.get("ocr") or {}).get("pages") or one.get("pages") or doc.get("pages")}
