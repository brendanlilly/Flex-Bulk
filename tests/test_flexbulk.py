"""Flex Bulk 1.0: the item, stock and price files load through the mapping; stock and prices come from them; orders
save, reopen and export; the staff app's endpoints match, price, look up and record choices."""
from __future__ import annotations

import io
import re
import json

import pandas as pd
import pytest

from flexbulk import loader


# ------------------------------------------------------------------ loading files
ITEMS = pd.DataFrame({
    "ITEM NO": ["1001", "1002", "1002", "1003", "", "1004"],
    "Item Description": ["NIBCO 3/4 in. Copper Sweat Ball Valve Full Port", "Charlotte 2 in. PVC Sch 40 90 Elbow",
                         "duplicate row", "Uponor 1/2 in. PEX-B Tubing 100 ft Coil", "no number", ""],
    "Vendor Name": ["NIBCO", "Charlotte", "x", "Uponor", "", "Acme"],
    "MPN": ["T-585-70-66", "PVC 02300 0600", "", "F1040500", "", "Z-1"],
    "Department": ["Plumbing", "Plumbing", "", "Plumbing", "", ""],
    "Status": ["Active", "Active", "", "Discontinued", "", ""],
})


def _cfg(layout="long"):
    c = loader.config()
    c["stock"]["layout"] = layout
    return c


@pytest.fixture()
def files(tmp_path):
    inc = tmp_path / "incoming"
    inc.mkdir()
    ITEMS.to_excel(inc / "items.xlsx", index=False)
    pd.DataFrame({"Item #": ["1001", "1001", "1002", "9999"], "Whse": ["100", "102", "100", "100"],
                  "QOH": ["12", "3", "0", "5"]}).to_csv(inc / "stock.csv", index=False)
    pd.DataFrame({"Item": ["1001", "1002", "1003", "7777"], "List Price": ["$1,204.50", "4.10", "", "9"],
                  "Contractor Price": ["800", "", "", "5"]}).to_parquet(inc / "prices.parquet", index=False)
    return inc


def test_items_stock_and_prices_load_through_the_mapping(files, tmp_path):
    out = tmp_path / "data"
    r = loader.load(log=lambda *_: None, incoming=files, data_dir=out, cfg=_cfg())
    p = pd.read_parquet(out / "canonical" / "products.parquet").set_index("sku")
    assert list(p.index) == ["1001", "1002", "1003"]                      # blank #, blank description, duplicate: out
    assert p.loc["1001", "brand"] == "NIBCO" and p.loc["1001", "mfr_part"] == "T-585-70-66"
    assert p.loc["1001", "list_price"] == pytest.approx(1204.50)          # "$1,204.50" read as a number
    assert p.loc["1003", "status"] == "D" and p.loc["1003", "list_price"] == 0
    assert json.loads(p.loc["1001", "attrs"]) == {"size": "3/4", "material": "copper", "connection": "sweat"}
    assert p.loc["1001", "cat3"] == "Plumbing" and p.loc["1001", "family"] == "plumbing"   # categories fall back
    assert p.loc["1001", "stock_tier"] != "dc" and p.loc["1002", "stock_tier"] == "dc"     # from the branches holding it
    st = pd.read_parquet(out / "canonical" / "stock.parquet")
    assert set(map(tuple, st[["sku", "branch_id"]].to_numpy())) == {("1001", "100"), ("1001", "102")}
    pr = pd.read_parquet(out / "canonical" / "prices.parquet").set_index("sku")
    assert pr.loc["1001", "Contractor"] == 800 and pd.isna(pr.loc["1002", "Contractor"])
    w = " ".join(r["warnings"])
    assert "duplicate item numbers" in w and "no item number" in w and "not on the item list" in w and "no list price" in w
    assert r["counts"]["items"] == 3 and r["price_levels"] == ["Contractor"]
    for t in ("accounts", "orders", "order_lines", "invoices", "content"):         # empty, but readable
        assert pd.read_parquet(out / "canonical" / f"{t}.parquet").empty


def test_wide_stock_layout(files, tmp_path):
    (files / "stock.csv").unlink()
    pd.DataFrame({"Item #": ["1001", "1003"], "QOH_100": [4, 0], "QOH_131": [0, 7]}).to_csv(files / "stock.csv", index=False)
    out = tmp_path / "data"
    loader.load(log=lambda *_: None, incoming=files, data_dir=out, cfg=_cfg("wide"))
    st = pd.read_parquet(out / "canonical" / "stock.parquet")
    assert sorted(map(tuple, st[["sku", "branch_id", "qty"]].to_numpy().tolist())) == [("1001", "100", 4.0), ("1003", "131", 7.0)]


def test_a_missing_required_column_says_which_and_what_the_file_has(files, tmp_path):
    ITEMS.rename(columns={"Item Description": "Words"}).to_excel(files / "items.xlsx", index=False)
    c = _cfg()
    c["items"].pop("compose")
    with pytest.raises(ValueError, match=r"no description column.*The file has: ITEM NO, Words"):
        loader.load(log=lambda *_: None, incoming=files, data_dir=tmp_path / "d", cfg=c, write=False)
    c["items"]["columns"]["sku"]["src"] = ["Nope"]
    with pytest.raises(ValueError, match=r"no column for sku"):
        loader.load(log=lambda *_: None, incoming=files, data_dir=tmp_path / "d", cfg=c, write=False)


def test_prices_fall_back_to_list(tmp_path):
    from flexbulk.prices import Prices
    pd.DataFrame({"sku": ["a", "b"], "list_price": [10.0, 0.0], "Net": [7.5, None]}).to_parquet(tmp_path / "p.parquet")
    pr = Prices(tmp_path / "p.parquet")
    assert pr.levels == ["Net"] and pr.price("a", "Net") == (7.5, "Net") and pr.price("a", None) == (10.0, "list")
    assert pr.price("b", "Net") == (None, "none")
    c = {"kind": "catalog", "sku": "a", "price": 99}
    pr.apply(c, "Net")
    assert c["price"] == 7.5 and c["list_price"] == 10.0 and c["price_basis"] == "Net"


def test_stock_comes_from_the_stock_file(engine):
    import numpy as np
    from webbsearch.config import settings
    av = engine.availability
    assert av.table is not None
    st = pd.read_parquet(settings.canonical_dir / "stock.parquet").head(50)
    for sku, b, q in zip(st["sku"], st["branch_id"], st["qty"]):
        i, bi = engine.catalog.sku_to_idx[sku], av.pos[str(b)]
        if av.counter[bi]:
            assert av._qty(np.array([i]), np.array([bi]))[0, 0] == int(round(q))


# ------------------------------------------------------------------ orders
def test_orders_save_reopen_and_export(tmp_path):
    from flexbulk.orders import Orders, to_csv, to_entry_csv, to_xlsx
    o = Orders(tmp_path / "o.sqlite")
    item = {"sku": "1001", "title": "NIBCO 3/4 Ball Valve", "brand": "NIBCO", "mfr_part": "T-585", "price": 12.5, "units": 4,
            "fulfil": {"text": "12 on hand at Bedford", "status": "pickup"}, "why": [{"k": "webb", "t": "Webb's #1 brand"}]}
    rows = [{"ln": {"n": 1, "text": "3/4 ball valve", "qty": 4, "crop": "data:image/png;base64,AAAA"}, "state": "ok",
             "item": item, "m": {"note": "", "options": [item], "crop": "x"}, "confirmed": True},
            {"ln": {"n": 2, "text": "=HYPERLINK(\"x\")", "qty": 1}, "state": "none", "item": None, "m": {"note": "Nothing relevant found"}},
            {"ln": {"n": 3, "text": "", "qty": 1}, "state": "new"}]
    s = o.save({"customer_no": "C-77", "customer_name": "Merrimack Mechanical", "job": "Library", "staff": "Dana", "rows": rows})
    assert s["id"].startswith("B") and s["n_lines"] == 2 and s["n_confirmed"] == 1 and s["total"] == 50.0
    got = o.get(s["id"])
    assert got["rows"][0]["ln"].get("crop") is None and "crop" not in got["rows"][0]["m"]     # bulky bits not kept
    o.save({**got, "job": "Library phase 2", "rows": got["rows"]})                             # same id: replaced
    assert o.get(s["id"])["job"] == "Library phase 2" and len(o.list()) == 1
    assert o.list(q="merrimack")[0]["id"] == s["id"] and o.customers("C-7")[0]["customer_name"] == "Merrimack Mechanical"
    csv_ = to_csv(o.get(s["id"])).decode("utf-8-sig")
    assert "C-77 · Merrimack Mechanical" in csv_ and "'=HYPERLINK" in csv_ and "Not found" in csv_
    assert to_entry_csv(o.get(s["id"])).decode("utf-8-sig").splitlines() == ["Item #,Qty", "1001,4"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(to_xlsx(o.get(s["id"]))))
    assert wb.sheetnames == ["Order", "Order entry"] and wb["Order entry"]["A2"].value == "1001"
    assert o.delete(s["id"]) and o.get(s["id"]) is None


# ------------------------------------------------------------------ the app
@pytest.fixture(scope="module")
def client(engine, tmp_path_factory):
    import os
    os.environ["FLEXBULK_ORDERS_DB"] = str(tmp_path_factory.mktemp("orders") / "o.sqlite")
    from fastapi.testclient import TestClient
    from webbsearch.api import server
    server._engine, server._bom, server._prices = engine, None, None
    import flexbulk.orders as fo
    fo._store = None
    return TestClient(server.app)


def test_setup_and_health(client):
    s = client.get("/api/setup").json()
    assert s["items"] > 1000 and s["branches"] and s["stock"] and s["levels"]
    assert client.get("/api/health").json()["app"] == "flex-bulk"


def test_match_prices_at_the_orders_level(client):
    lv = client.get("/api/setup").json()["levels"][0]
    body = {"lines": [{"text": "3/4 ball valve sweat", "qty": 6}, {"text": "", "qty": 1}], "branch": "100"}
    a = client.post("/api/bom/match", json=body).json()
    b = client.post("/api/bom/match", json={**body, "level": lv}).json()
    assert len(a["lines"]) == 1 and a["level"] == "List" and b["level"] == lv
    oa, ob = a["lines"][0]["options"][0], b["lines"][0]["options"][0]
    assert oa["sku"] == ob["sku"] and oa["price"] == oa["list_price"] and ob["price"] < oa["price"]
    assert oa["fulfil"]["status"] in ("pickup", "partial", "transfer", "dc", "special")


def test_lookup_and_item_detail(client):
    d = client.get("/api/lookup", params={"q": "pex b tubing 1/2", "branch": "100", "qty": 200, "uom": "FT"}).json()
    assert d["items"] and all("fulfil" in c for c in d["items"])
    sku = d["items"][0]["sku"]
    p = client.get(f"/api/product/{sku}", params={"branch": "100"}).json()
    assert p["sku"] == sku and p["branches"] and p["branches"][0]["home"]
    assert client.get("/api/product/not-an-item").status_code == 404


def test_a_customers_own_picks_count_for_more(client, engine):
    line = {"text": "3/4 sweat ball valve full port", "qty": 4}
    before = client.post("/api/bom/match", json={"lines": [line], "branch": "100", "customer": "C-1"}).json()["lines"][0]
    pick = before["options"][3]
    for i in range(1):            # one pick: enough for this customer, not yet for everyone (learned_min = 2)
        r = client.post("/api/track", json={"kind": "bulk_choice", "sid": f"s{i}", "sku": pick["sku"], "customer": "C-1",
                                            "meta": {"key": before["key"], "brand": pick["brand"], "family": pick["family"]}})
        assert r.json()["ok"]
    mine = client.post("/api/bom/match", json={"lines": [line], "branch": "100", "customer": "C-1"}).json()["lines"][0]
    other = client.post("/api/bom/match", json={"lines": [line], "branch": "100", "customer": "C-2"}).json()["lines"][0]
    pos_mine = [o["sku"] for o in mine["options"]].index(pick["sku"])
    assert pos_mine < 3 and any(w["k"] == "learned" for w in mine["options"][pos_mine]["why"])
    assert [o["sku"] for o in other["options"]].index(pick["sku"]) == 3          # another customer: unchanged
    assert client.post("/api/track", json={"kind": "click", "sku": "x"}).status_code == 422


def test_an_item_cross_to_an_item_not_on_the_list_is_ignored(engine):
    from webbsearch.api import server
    from webbsearch.retrieval.entitlements import ANONYMOUS
    s = {**engine.merch.get(), "crosses": [{"id": "x1", "kind": "item", "enabled": True, "scope": "company",
                                             "from_text": "kohler highline toilet", "from_part": "", "from_brand": "",
                                             "to_sku": "NOT-ON-LIST", "note": ""}]}
    m = server.bom().match({"text": "kohler highline toilet", "qty": 1}, ANONYMOUS, None, settings=s)
    assert m["status"] in ("rec", "none") and "cross" not in m


def test_orders_through_the_api(client):
    r = client.post("/api/orders", json={"customer_no": "C-9", "customer_name": "Kennebec Plumbing", "staff": "Pat",
                                         "rows": [{"ln": {"text": "3/4 ball valve", "qty": 2}, "state": "ok",
                                                   "item": {"sku": "1", "title": "Valve", "price": 3.0, "units": 2}}]}).json()
    assert r["n_confirmed"] == 1 and r["total"] == 6.0
    assert client.get("/api/orders", params={"q": "kennebec"}).json()[0]["id"] == r["id"]
    x = client.get(f"/api/orders/{r['id']}/export", params={"fmt": "xlsx"})
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Kennebec-Plumbing" in x.headers["content-disposition"]
    assert client.get(f"/api/orders/{r['id']}/export", params={"fmt": "pdf"}).status_code == 422
    assert client.delete(f"/api/orders/{r['id']}").json()["ok"]
    assert client.get(f"/api/orders/{r['id']}").status_code == 404


def test_the_page_and_dashboard_load(client):
    h = client.get("/").text
    assert "Flex<b>Bulk</b>" in h and "/static/bulk.js?v=" in h and "/static/app.js" not in h
    a = client.get("/admin").text
    assert "Flex<b>Bulk</b> <span class=\"w2\">Settings</span>" in a and 'data-page="crosses"' in a and 'data-page="overview"' not in a


def test_an_attribute_export_gets_descriptions_built(tmp_path):
    """Webb's PIM export: pipe-delimited Windows-1252 text, no description column, one column per attribute."""
    inc = tmp_path / "incoming"
    inc.mkdir()
    head = "Control #|Manufacturer Part Number|Manufacturer Name|Modifier Noun|Product Code|Size|Material|End Connection|Port Type|Product Type|Finish|Bend Angle"
    rows = ["51414|51414|Webstone|Ball Valve|H - Heating|1 IN|Forged Brass|SweatxRotating Flange Drain|Full Port|||",
            "526401|526401|Georg Fischer|Elbow|L - Plastics|1-1/2 IN|Polypropylene|SocketxSpigot||Street||45 DEG",
            "9|X1|Miscellaneous|Tee|C - Commodities|||||||",
            "10|Z|Acme||N - Non-Product|||||||",
            "11|D 1/2\"|Décor Co|Pull|D - Décor|3 IN|Brass|||Bar|Satin Brass|"]
    (inc / "items.txt").write_bytes(("\n".join([head] + rows) + "\n").encode("cp1252"))
    r = loader.load(log=lambda *_: None, incoming=inc, data_dir=tmp_path / "d", cfg=loader.config())
    p = pd.read_parquet(tmp_path / "d" / "canonical" / "products.parquet").set_index("sku")
    assert p.loc["51414", "title"] == "Webstone 1 in. Forged Brass Sweat x Rotating Flange Drain Ball Valve, Full Port"
    assert p.loc["526401", "title"] == "Georg Fischer 1-1/2 in. Polypropylene Socket x Spigot Elbow, Street, 45 DEG"
    assert p.loc["9", "title"] == "Tee" and p.loc["9", "brand"] == ""                       # "Miscellaneous": no brand
    assert "10" not in p.index and any("not products" in w for w in r["warnings"])
    assert p.loc["11", "brand"] == "Décor Co" and p.loc["11", "mfr_part"] == 'D 1/2"'      # cp1252 and an inch mark
    a = json.loads(p.loc["51414", "attrs"])
    assert a["size"] == "1" and a["Port Type"] == "Full Port" and a["joint"] == "SxR"
    assert json.loads(p.loc["526401", "attrs"])["joint"] == "SxS" and "Port Type: Full Port" in p.loc["51414", "description"]
    assert p.loc["51414", "cat1"] == "Heating" and p.loc["51414", "cat3"] == "Ball Valve" and p.loc["51414", "family"] == "ball_valve"


def test_a_manufacturers_number_part_with_the_brand_on_the_line_is_exact(client, engine):
    import numpy as np
    cat = engine.catalog
    j = next(i for i in range(cat.n) if str(cat.mfr_part[i]).isdigit() and len(str(cat.mfr_part[i])) >= 4 and cat.brand[i]
             and len(cat.id_lookup(str(cat.mfr_part[i]))) == 1)
    x = client.post("/api/bom/match", json={"lines": [{"text": f"{cat.brand[j]} {cat.mfr_part[j]}", "qty": 1}]}).json()["lines"][0]
    assert x["status"] == "exact" and x["options"][0]["sku"] == cat.sku[j]


def test_every_match_links_to_its_web_store_page(client):
    x = client.post("/api/bom/match", json={"lines": [{"text": "3/4 ball valve", "qty": 1}]}).json()["lines"][0]
    o = x["options"][0]
    assert o["web_url"] == f"https://www.fwwebb.com/product/{o['sku']}"
    assert client.get(f"/api/product/{o['sku']}").json()["web_url"] == o["web_url"]


def test_match_reports_from_orders_log_and_a_test_file(client, tmp_path, monkeypatch):
    import flexbulk.report as rp
    from openpyxl import load_workbook
    monkeypatch.setattr(rp, "REPORTS", tmp_path / "reports")
    m = client.post("/api/bom/match", json={"lines": [{"text": "3/4 ball valve sweat", "qty": 2}, {"text": "1/2 pex b tubing", "qty": 1},
                                                      {"text": "zzqx flux capacitor", "qty": 1}], "sid": "rep1"}).json()["lines"]
    second = m[0]["options"][1]
    rows = [{"ln": {"n": 1, "text": "3/4 ball valve sweat", "qty": 2}, "m": m[0], "ch": {"t": "opt", "i": 1}, "state": "ok", "item": second, "confirmed": True},
            {"ln": {"n": 2, "text": "1/2 pex b tubing", "qty": 1}, "m": m[1], "ch": {"t": "opt", "i": 0}, "state": "ok", "item": m[1]["options"][0], "confirmed": True},
            {"ln": {"n": 3, "text": "zzqx flux capacitor", "qty": 1}, "m": m[2], "state": "none"}]
    client.post("/api/orders", json={"customer_no": "C-5", "staff": "Pat", "rows": rows})
    r = rp.make("orders", 30)
    s = r["summary"]
    assert s["lines"] == 3 and s["with_expected"] == 2 and s["top1"] == 0.5 and s["top3"] == 1.0
    assert s["how_staff_chose"] == {"chose another option": 1, "kept top pick": 1}
    wb = load_workbook(rp.REPORTS / f"{r['meta']['id']}.xlsx")
    assert wb.sheetnames == ["Summary", "Lines", "Misses", "Trace", "Not found", "Brands", "Changed", "Settings", "About"]
    miss = list(wb["Misses"].iter_rows(min_row=2, values_only=True))
    assert len(miss) == 1 and miss[0][0] == "3/4 ball valve sweat" and miss[0][2] == second["sku"] and miss[0][4] == 2
    assert all(v != "'" for row in wb["Lines"].iter_rows(values_only=True) for v in row)          # no stray apostrophes
    # a second run from the same source is compared with the first
    r2 = rp.make("orders", 30)
    assert r2["meta"]["previous"]["id"] == r["meta"]["id"] and r2["meta"]["previous"]["top1"] == 0.5
    # the app's log: a choice made on a matched line is its answer
    client.post("/api/track", json={"kind": "bulk_choice", "sid": "rep1", "sku": second["sku"], "meta": {"key": m[0]["key"], "how": "chosen"}})
    lg = rp.make("log", 1)["summary"]
    assert lg["lines"] >= 3 and lg["with_expected"] >= 1
    # a test file
    f = "Line,Expected,Qty\n3/4 ball valve sweat,%s,2\n1/2 pex b tubing,,1\n" % m[0]["options"][0]["sku"]
    t = rp.make("file", file_name="t.csv", file_data=f.encode())["summary"]
    assert t["lines"] == 2 and t["with_expected"] == 1 and t["top1"] == 1.0
    with pytest.raises(ValueError, match="no line column"):
        rp.make("file", file_name="t.csv", file_data=b"Foo,Bar\n1,2\n")


def test_reports_through_the_dashboard(client):
    import time as _t
    from webbsearch.api.admin import TOKEN
    h = {"X-Admin-Token": TOKEN}
    assert client.post("/api/admin/reports", json={"source": "orders"}).status_code == 403
    assert client.post("/api/admin/reports", json={"source": "orders", "days": 30}, headers=h).json()["ok"]
    for _ in range(100):
        d = client.get("/api/admin/reports", headers=h).json()
        if not d["job"]["running"]:
            break
        _t.sleep(0.2)
    assert d["job"]["error"] is None and d["reports"]
    x = client.get(f"/api/admin/reports/{d['reports'][0]['id']}.xlsx", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_why_a_line_matched_what_it_did(client):
    x = client.post("/api/bom/explain", json={"line": {"text": "3/4 cxc 90 elbow", "qty": 4}, "branch": "100"}).json()
    steps = [t["step"] for t in x["trace"]]
    assert steps[:4] == ["read", "crosses", "exact lookup", "description"] and "ranked" in steps and steps[-1] == "result"
    rk = next(t for t in x["trace"] if t["step"] == "ranked")
    top = rk["candidates"][0]
    assert top["sku"] == x["options"][0]["sku"] and top["product"]
    assert {"search", "fit", "bonus", "bonus_parts", "product", "size", "spec_ok"} <= set(top)
    assert not any("_parts" in o or "_trace" in o for o in x["options"])
    assert "trace" not in client.post("/api/bom/match", json={"lines": [{"text": "3/4 cxc 90 elbow", "qty": 4}]}).json()["lines"][0]


def test_a_test_case_is_read_again_and_scored(client, tmp_path, monkeypatch):
    import base64
    import zipfile
    import flexbulk.report as rp
    from pathlib import Path
    from openpyxl import load_workbook
    monkeypatch.setattr(rp, "REPORTS", tmp_path / "reports")
    raw = (Path(__file__).resolve().parents[1] / "samples" / "material-list.docx").read_bytes()
    b64 = base64.b64encode(raw).decode()
    doc = client.post("/api/bom/parse", json={"filename": "material-list.docx", "data_b64": b64}).json()
    lines = doc["lines"]
    m = client.post("/api/bom/match", json={"lines": lines}).json()["lines"]
    rows = [{"ln": ln, "m": x, "ch": {"t": "opt", "i": 0}, "state": "ok" if x["options"] else "none",
             "item": x["options"][0] if x["options"] else None, "confirmed": True} for ln, x in zip(lines, m)]
    rows[0]["ln"] = {**rows[0]["ln"], "qty": 999}                    # staff "corrected" a quantity: reading should flag it
    o = client.post("/api/orders", json={"customer_no": "T-1", "staff": "Pat", "test": True, "rows": rows,
                                         "files": [{"name": "material-list.docx", "b64": b64}]}).json()
    assert o["test"] == 1 and client.get(f"/api/orders/{o['id']}").json()["files"] == ["material-list.docx"]
    assert [x["id"] for x in client.get("/api/orders", params={"tests": True}).json()] == [o["id"]]
    r = rp.make("tests")
    s = r["summary"]
    answered = sum(1 for x in m if x["options"])
    assert s["with_expected"] == answered and s["top1"] == 1.0
    rd = json.loads((rp.REPORTS / f"{r['meta']['id']}.json").read_text())["summary"]["reading"]
    assert rd["lines_read"] == len(lines) and rd["checked"] == len(lines) and rd["read_right"] == len(lines) - 1
    wb = load_workbook(rp.REPORTS / f"{r['meta']['id']}.xlsx")
    assert {"Trace", "Reading"} <= set(wb.sheetnames)
    steps = [row[-1] for row in wb["Lines"].iter_rows(min_row=2, values_only=True)]
    assert all(st and st.startswith("read:") and "result:" in st for st in steps)
    checks = [row[-1] for row in wb["Reading"].iter_rows(min_row=2, values_only=True) if row[7] and row[7].startswith("item")]
    assert sum(1 for c in checks if c.startswith("qty: read")) == 1
    z = zipfile.ZipFile(io.BytesIO(rp.bundle(r["meta"]["id"])))
    names = z.namelist()
    assert {"README.md", "report.xlsx", "settings.json", "code/webbsearch/retrieval/bom.py", "code/webbsearch/api/static/bulk.js"} <= set(names)
    assert any(n.startswith("inputs/") and n.endswith("material-list.docx") for n in names)
    assert not any(n.startswith("inputs/") for n in zipfile.ZipFile(io.BytesIO(rp.bundle(r["meta"]["id"], inputs=False))).namelist())
    client.delete(f"/api/orders/{o['id']}")


def test_customers_documents_as_a_report(client, tmp_path, monkeypatch):
    import flexbulk.report as rp
    from pathlib import Path
    monkeypatch.setattr(rp, "REPORTS", tmp_path / "reports")
    raw = (Path(__file__).resolve().parents[1] / "samples" / "order-email.eml").read_bytes()
    r = rp.make("documents", documents=[("order-email.eml", raw)])
    assert r["summary"]["lines"] >= 6 and r["summary"]["with_expected"] == 0
    rd = json.loads((rp.REPORTS / f"{r['meta']['id']}.json").read_text())["summary"]["reading"]
    assert rd["lines_read"] >= 6 and rd["set_aside"] >= 5                          # greeting, signature, address...
    assert (rp.REPORTS / f"{r['meta']['id']}-inputs" / "order-email.eml").exists()


# ------------------------------------------------------------------ 1.2: files with no header, brands, shorthand
def test_a_table_with_no_header_row_is_read_by_its_columns():
    from webbsearch.retrieval.resolve import clean_part, parse_table
    rows = ["Amtrol\tEX-30\tEX-30 AMTROL EXTROL TANK\tWater Heater / Tank Accessories\tEA",
            "3M / Venture Tape\t7010378552\t2 X 60YD XH BLK DUCT TAPE\tVenting, Flue & Ductwork\tEA",
            "Caleffi\t521516A\tMXG VLV ADJ 3W THM W/GA 3/4\" LL BRS PRS\tValves & Regulators\tEA",
            "Commodity / Distributor House Line\t1/2inch\t1/2inch BLACK MALLEABLE 90\tFittings & Adapters\tEA",
            "Aprilaire / Research Products\tSee description\t800 APRILAIRE STEAM HUMIDIFIER\tAir Filters & Indoor Air Quality\tEA",
            "Amtrol\tST-12\tST-12 AMTROL THERM-X-TROL TANK\tWater Heater / Tank Accessories\tEA"]
    t = parse_table("\n".join(rows))
    assert t["roles"] == {"uom": 4, "desc": 2, "part": 1, "brand": 0, "category": 3} and t["header_row"] == -1 and t["inferred"]
    L = t["lines"]
    assert len(L) == 6 and all(x["qty"] == 1 and x["uom"] == "EA" for x in L)
    assert L[0]["part"] == "EX-30" and L[0]["brand"] == "Amtrol" and "Water Heater" not in L[0]["text"] + L[0]["desc"]
    assert L[1]["part"] == "7010378552"
    assert L[3]["part"] == "" and L[3]["text"].endswith("1/2inch BLACK MALLEABLE 90")         # a size isn't a part #
    assert L[4]["part"] == ""                                                                 # "See description"
    assert clean_part("VCMA-20ULS-PRO,") == ("VCMA-20ULS-PRO", "") and clean_part("3/8X3/4X25'") == ("", "3/8X3/4X25'")
    # a list with a header still reads by its header
    assert parse_table("Qty\tDescription\n4\t3/4 ball valve\n2\t1/2 elbow")["header_row"] == 0


def test_brand_names_resolve_vendor_spellings(tmp_path):
    from flexbulk.brands import BrandNames
    bn = BrandNames(["Mueller", "A.O. Smith", "State", "Honeywell Home", "Legend", "Quick-Sling", "Diversitech", "GE"],
                    path=tmp_path / "b.json")
    assert bn.resolve("Mueller Industries") == {"Mueller"}
    assert bn.resolve("A. O. Smith / State") == {"A.O. Smith", "State"}
    assert bn.resolve("Resideo (Honeywell Home)") == {"Honeywell Home"}
    assert bn.resolve("Legend Valve") == {"Legend"} and bn.resolve("Acme Widgets") == frozenset()
    assert bn.same("Diversitech", "Quick-Sling") is False
    bn.add("Diversitech", ["Quick-Sling"])
    assert bn.resolve("Diversitech") == {"Diversitech", "Quick-Sling"}
    from collections import Counter
    bn.suggest(Counter({("Bindus", "GE"): 2, ("Hydrolevel", "GE"): 1}))
    assert [s["name"] for s in bn.table()["suggested"]] == ["Bindus"]
    bn.dismiss("Bindus", "GE")
    bn.suggest(Counter({("Bindus", "GE"): 5}))
    assert bn.table()["suggested"] == []


def test_shorthand_expands_and_learns(tmp_path):
    from flexbulk.shorthand import Shorthand
    sh = Shorthand(path=tmp_path / "s.json")
    assert sh.expand('MXG VLV ADJ 3W THM 3/4" BRS') == 'mixing valve adjustable 3W thermostatic 3/4" brass'
    assert sh.expand("ST-12 AMTROL THERM-X-TROL") == "ST-12 AMTROL THERM-X-TROL"          # part #s and names untouched
    assert sh.expand("11/4 X 10 M COPPER 1/2inch") == '1-1/4 X 10 M COPPER 1/2"'
    pairs = [("ANTFRZ CRYO-TEK 100 5GAL", "Antifreeze Cryo-Tek 5 GAL Heat Transfer Fluid")] * 3 + \
            [("HMDFR PAD 35", "Humidifier Pad Water Panel")]
    got = sh.learn(pairs, known_words=lambda w: w in {"pad"})
    assert got["on"] == {"antfrz": "antifreeze"} and got["review"] == {"hmdfr": "humidifier"}
    assert sh.expand("ANTFRZ 5GAL") == "antifreeze 5GAL"
    sh.switch_off("antfrz")
    assert sh.expand("ANTFRZ") == "ANTFRZ" and sh.listing()["off"] == ["antfrz"]


def test_a_part_number_of_another_brand_is_offered_not_confirmed(client, engine):
    cat = engine.catalog
    j = next(i for i in range(cat.n) if len(str(cat.mfr_part[i])) >= 6 and cat.brand[i] and len(cat.id_lookup(
        __import__("webbsearch.retrieval.analyzer", fromlist=["compact_id"]).compact_id(str(cat.mfr_part[i])))) == 1)
    other = next(b for b in set(map(str, cat.brand)) if b and b != cat.brand[j] and len(b) >= 5 and " " not in b)
    words = " ".join(str(cat.title[j]).split()[1:4])
    right = client.post("/api/bom/match", json={"lines": [{"text": f"{cat.brand[j]} {cat.mfr_part[j]}", "brand": str(cat.brand[j]),
                                                           "part": str(cat.mfr_part[j]), "desc": words, "qty": 1}]}).json()["lines"][0]
    assert right["status"] == "exact" and right["confirmed"]
    x = client.post("/api/bom/match", json={"lines": [{"text": f"{other} {cat.mfr_part[j]}", "brand": other, "part": str(cat.mfr_part[j]),
                                                       "desc": words, "qty": 1}]}).json()["lines"][0]
    assert x["status"] == "rec" and not x["confirmed"] and cat.sku[j] in [o["sku"] for o in x["options"]]
    assert f"is {cat.brand[j]}'s" in x["note"]


def test_a_close_part_number_of_the_lines_brand(client, engine):
    from webbsearch.retrieval.analyzer import compact_id
    cat = engine.catalog
    j = next(i for i in range(cat.n) if re.fullmatch(r"[A-Z]{2,}\d{3,}[A-Z0-9-]*", str(cat.mfr_part[i])) and len(compact_id(cat.mfr_part[i])) >= 7
             and cat.brand[i] and not cat.id_lookup(compact_id(cat.mfr_part[i])[:-2]))
    short = compact_id(cat.mfr_part[j])[:-2]
    x = client.post("/api/bom/match", json={"lines": [{"text": f"{cat.brand[j]} {short}", "brand": str(cat.brand[j]), "part": short,
                                                       "desc": "", "qty": 1}]}).json()["lines"][0]
    assert x["options"] and x["options"][0]["sku"] == cat.sku[j] and not x["confirmed"] and "Close part #" in x["note"]


def test_hide_the_part_numbers_report(client, engine, tmp_path, monkeypatch):
    import flexbulk.report as rp
    import flexbulk.shorthand as shm
    from openpyxl import Workbook
    monkeypatch.setattr(rp, "REPORTS", tmp_path / "reports")
    monkeypatch.setattr(shm, "_sh", shm.Shorthand(path=tmp_path / "sh.json"))
    cat = engine.catalog
    picks = [i for i in range(0, cat.n, 97) if len(str(cat.mfr_part[i])) >= 5 and cat.brand[i]][:40]
    wb = Workbook()
    ws = wb.active
    for i in picks:
        ws.append([f"{cat.brand[i]} Inc", str(cat.mfr_part[i]), f"{cat.mfr_part[i]} {' '.join(str(cat.title[i]).split()[1:])}", "Some Category", "EA"])
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    r = rp.make("documents", documents=[("vendor.xlsx", buf.getvalue())], hide=True)
    s = r["summary"]
    assert s["with_expected"] >= 30 and s["lines"] == s["with_expected"]
    assert s["top3"] >= 0.5                                    # titles minus the brand still find most of them
    assert r["meta"]["hide"] and r["meta"]["learned"]["pairs"] >= 30


def test_brand_names_and_shorthand_pages(client, engine, tmp_path, monkeypatch):
    import flexbulk.brands as bm
    import flexbulk.shorthand as shm
    monkeypatch.setattr(bm, "_bn", None)
    monkeypatch.setenv("FLEXBULK_BRAND_NAMES", str(tmp_path / "bn.json"))
    monkeypatch.setattr(shm, "_sh", shm.Shorthand(path=tmp_path / "sh.json"))
    from webbsearch.api.admin import TOKEN
    h = {"X-Admin-Token": TOKEN}
    d = client.get("/api/admin/names", headers=h).json()
    assert d["brands"]["names"] == [] and "vlv" in d["shorthand"]["starter"]
    b = next(x for x in set(map(str, engine.catalog.brand)) if len(x) >= 5)
    assert client.post("/api/admin/names/brand", json={"name": "Some Vendor Name", "brands": [b]}, headers=h).status_code == 200
    assert client.get("/api/admin/names", headers=h).json()["brands"]["names"][0]["brands"] == [b]
    assert client.post("/api/admin/names/brand", json={"name": "x", "brands": ["Not A Brand On The List"]}, headers=h).status_code == 422
    client.post("/api/admin/names/shorthand", json={"abbr": "QQV", "word": "quarterturn", "action": "accept"}, headers=h)
    assert shm.get().expand("QQV 1/2") == "quarterturn 1/2"
    client.post("/api/admin/names/brand/remove", json={"name": "Some Vendor Name"}, headers=h)
    assert client.get("/api/admin/names", headers=h).json()["brands"]["names"] == []
