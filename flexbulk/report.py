"""Match reports: run lines through the matcher and write one Excel workbook that says how well it did, line by
line, for working on the matching (drop the workbook into a conversation with Claude, or read it yourself).

Where the lines come from:
  tests      orders kept as test cases (the staff app's "Test case" box). Their customers' documents are read
             again, so a report scores the reading (lines, quantities) as well as the matching; a confirmed line's
             item is its answer.
  documents  customers' files given to the report (emails, Excel, Word, PDFs, photos): read and matched, no answers.
  orders   saved orders. A confirmed line's item is the answer the matcher should have given; lines left
           unconfirmed or not found are reported too.
  log      lines matched in the app (analytics), with what staff then chose for them when they chose.
  file     a test file: CSV / Excel with a line column ("Line", "Requested", "Description"...) and, optionally,
           the item # it should match ("Expected", "Item #") and a quantity.

Each line is matched again with today's settings, so the report measures the matcher as it is now, and says what
was shown at the time when it knows (orders, log).

    python -m flexbulk.report --source orders --days 90
    python -m flexbulk.report --source file --file my-test-lines.xlsx
    (or the dashboard: Match reports)

Workbook: Summary (the numbers, and the change since the previous report from the same source), Lines (every
line: what was asked, the top matches with their reasons, the expected item and where it ranked), Misses (expected
item not ranked first), Not found, Brands (lines naming brands not on the list), Changed (now differs from what was
shown then), Settings (the weights and crosses in force), About (what each column means).
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import secrets
import time
from collections import Counter, defaultdict
from pathlib import Path

from webbsearch.config import settings

REPORTS = Path(os.environ.get("FLEXBULK_REPORTS") or settings.data_dir / "bulk" / "reports")
K = 8
SOURCES = ("orders", "log", "file", "tests", "documents")


# ------------------------------------------------------------------ the lines
def from_orders(days: int, limit: int = 5000) -> list[dict]:
    from flexbulk.orders import store
    st = store()
    since = time.time() - days * 86400
    cases = []
    for o in st.list(limit=500):
        if o["updated"] < since:
            continue
        full = st.get(o["id"])
        for k, r in enumerate(full["rows"]):
            ln = r.get("ln") or {}
            if not (ln.get("text") or ln.get("desc")):
                continue
            m = r.get("m") or {}
            it = r.get("item") or {}
            opts = [x.get("sku") for x in (m.get("options") or [])]
            ok = r.get("state") == "ok" and not r.get("excluded") and it.get("sku")
            ch = r.get("ch") or {}
            how = ("found by hand" if m.get("via") == "Found by hand" else "substitute" if ch.get("t") == "sub"
                   else "kept top pick" if ok and it.get("sku") == (opts[0] if opts else None) else
                   "chose another option" if ok else "")
            cases.append({"src": f"order {o['id']} line {ln.get('n') or k + 1}", "customer": o.get("customer_no") or o.get("customer_name") or "",
                          "text": ln.get("text") or "", "desc": ln.get("desc") or "", "qty": ln.get("qty") or 1,
                          "uom": ln.get("uom") or "", "brand": ln.get("brand") or "", "part": ln.get("part") or "",
                          "from_ocr": bool(ln.get("from_ocr")), "expected": it.get("sku") if ok else None, "how": how,
                          "then_status": m.get("status") or "", "then_top": opts[0] if opts else None,
                          "state": "left out" if r.get("excluded") else r.get("state") or ""})
            if len(cases) >= limit:
                return cases
    return cases


def from_log(days: int, limit: int = 5000) -> list[dict]:
    from webbsearch.api.admin import analytics
    a = analytics()
    if a is None:
        return []
    since = time.strftime("%Y-%m-%d", time.localtime(time.time() - days * 86400))
    choices = {}
    for sess, acct, sku, meta in a._q("SELECT session, account, sku, meta FROM events WHERE kind='bulk_choice' AND day>=? ORDER BY id",
                                      (since,)):
        d = json.loads(meta or "{}")
        if d.get("key"):
            choices[(sess, d["key"])] = (sku, d.get("how") or "", d.get("rank"))
    cases, seen = [], set()
    for sess, acct, qty, sku, meta, day in a._q("SELECT session, account, qty, sku, meta, day FROM events WHERE kind='bulk_line' "
                                                "AND day>=? ORDER BY id DESC", (since,)):
        d = json.loads(meta or "{}")
        text = d.get("text") or ""
        if not text or (sess, d.get("key"), text) in seen:
            continue
        seen.add((sess, d.get("key"), text))
        ch = choices.get((sess, d.get("key")))
        cases.append({"src": f"logged {day}", "customer": (acct or "").removeprefix("cust:"), "text": text, "desc": "",
                      "qty": qty or 1, "uom": "", "brand": "", "part": "", "from_ocr": d.get("source") == "handwriting/scan",
                      "expected": ch[0] if ch else None,
                      "how": ({"confirm": "kept top pick", "confirm top pick": "kept top pick", "chosen": "chose another option",
                               "substitute": "substitute", "search": "found by hand", "nearest": "took nearest"}.get(ch[1], ch[1]) if ch else ""),
                      "then_status": d.get("status") or "", "then_top": sku, "state": ""})
        if len(cases) >= limit:
            break
    return cases[::-1]


def from_file(name: str, data: bytes, limit: int = 5000) -> list[dict]:
    """A test file: a line column and, optionally, the item # it should match."""
    from flexbulk.loader import pick, read_any
    import tempfile
    suffix = Path(name).suffix or ".csv"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as f:
        f.write(data)
        tmp = Path(f.name)
    try:
        df = read_any(tmp)
    finally:
        tmp.unlink(missing_ok=True)
    line = pick(df, ["Line", "Requested", "Text", "Query", "Description", "Item Description", "Search"])
    if line is None:
        raise ValueError(f"no line column (Line, Requested, Text, Description...). The file has: {', '.join(map(str, df.columns[:20]))}")
    exp = pick(df, ["Expected", "Expected Item", "Expected Item #", "Answer", "Item #", "Control #", "SKU"])
    qty = pick(df, ["Qty", "Quantity", "QTY"])
    uom = pick(df, ["UOM", "U/M", "Unit"])
    brand = pick(df, ["Manufacturer", "Brand", "Mfr"])
    cases = []
    for k, r in df.head(limit).iterrows():
        t = str(r[line]).strip()
        if not t:
            continue
        q = re.sub(r"[^\d.]", "", str(r[qty])) if qty else ""
        cases.append({"src": f"{name} row {k + 2}", "customer": "", "text": t, "desc": "", "qty": int(float(q)) if q else 1,
                      "uom": str(r[uom]).strip() if uom else "", "brand": str(r[brand]).strip() if brand else "", "part": "",
                      "from_ocr": False, "expected": (str(r[exp]).strip() or None) if exp else None, "how": "",
                      "then_status": "", "then_top": None, "state": ""})
    return cases


def read_doc(name: str, data: bytes) -> dict:
    """A customer's document read the way the app reads it (bom.parse_request: email, Excel, Word, PDF, OCR)."""
    import base64
    from webbsearch.api.server import engine
    from webbsearch.retrieval.bom import parse_request
    e = engine()
    vocab = getattr(getattr(e.retriever, "b", None), "term_exists", None)
    return parse_request(None, name, base64.b64encode(data).decode(), None, None, None, vocab=vocab, claude=e.claude)


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", str(t or "").lower()).strip()


def doc_cases(name: str, doc: dict, reading: list, label: str, saved: dict | None = None, customer: str = "") -> list[dict]:
    """Lines read from one document, as cases; `saved` (a test case's rows by (source, line n) and by text) gives the
    answer and what the line should have been read as."""
    cases = []
    for s_ in doc.get("sources") or []:
        for o in s_.get("other") or []:
            reading.append({"doc": name, "source": s_.get("label"), "n": o.get("n"), "raw": o.get("raw"), "text": o.get("text"),
                            "qty": o.get("qty"), "uom": o.get("uom"), "read_as": "set aside", "why": o.get("why"), "check": ""})
    used = set()
    for ln in doc.get("lines") or []:
        sv = None
        if saved is not None:
            sv = saved.get(("n", ln.get("src_label"), ln.get("n"))) or saved.get(("t", _norm(ln.get("text"))))
            if sv is not None:
                used.add(id(sv))
        check = ""
        if saved is not None:
            if sv is None:
                check = "not in the test case"
            else:
                want = sv.get("ln") or {}
                diffs = []
                if _norm(want.get("text")) != _norm(ln.get("text")):
                    diffs.append(f"text: read “{ln.get('text')}”, should be “{want.get('text')}”")
                if int(want.get("qty") or 1) != int(ln.get("qty") or 1):
                    diffs.append(f"qty: read {ln.get('qty')}, should be {want.get('qty')}")
                check = "; ".join(diffs) or "ok"
        reading.append({"doc": name, "source": ln.get("src_label"), "n": ln.get("n"), "raw": ln.get("raw"), "text": ln.get("text"),
                        "qty": ln.get("qty"), "uom": ln.get("uom"), "read_as": "item line" + (" (handwriting)" if ln.get("from_ocr") else ""),
                        "why": f"confidence {ln['conf']:.2f}" if isinstance(ln.get("conf"), (int, float)) else "", "check": check})
        it = (sv or {}).get("item") or {}
        ok = sv is not None and sv.get("state") == "ok" and not sv.get("excluded") and it.get("sku")
        m = (sv or {}).get("m") or {}
        opts = [x.get("sku") for x in (m.get("options") or [])]
        cases.append({"src": f"{label} · {ln.get('src_label') or name} line {ln.get('n')}", "customer": customer,
                      "text": ln.get("text") or "", "desc": ln.get("desc") or "", "qty": ln.get("qty") or 1, "uom": ln.get("uom") or "",
                      "brand": ln.get("brand") or "", "part": ln.get("part") or "", "from_ocr": bool(ln.get("from_ocr")),
                      "expected": it.get("sku") if ok else None, "how": "", "then_status": m.get("status") or "",
                      "then_top": opts[0] if opts else None, "state": (sv or {}).get("state") or "", "reading": check})
    if saved is not None:
        for sv in {id(v): v for v in saved.values()}.values():
            if id(sv) not in used and (sv.get("ln") or {}).get("src_label"):
                ln = sv.get("ln") or {}
                reading.append({"doc": name, "source": ln.get("src_label"), "n": ln.get("n"), "raw": ln.get("raw"), "text": ln.get("text"),
                                "qty": ln.get("qty"), "uom": ln.get("uom"), "read_as": "NOT READ", "why": "", "check": "in the test case, not read from the document now"})
    return cases


def from_tests(reading: list, limit: int = 5000) -> tuple[list[dict], list[str]]:
    from flexbulk.orders import store
    st = store()
    cases, inputs = [], []
    for o in st.list(limit=500, tests=True):
        full = st.get(o["id"])
        files = st.files(o["id"])
        cust = o.get("customer_no") or o.get("customer_name") or ""
        if not files:                                # typed or pasted: the saved lines are the input
            for c in from_orders_rows(full):
                cases.append(c)
            continue
        by_src = defaultdict(dict)
        for r in full["rows"]:
            ln = r.get("ln") or {}
            key_src = ln.get("src_label") or ""
            by_src[key_src][("n", key_src, ln.get("n"))] = r
            by_src[key_src][("t", _norm(ln.get("raw") or ln.get("text")))] = r
            by_src[key_src][("t", _norm(ln.get("text")))] = r
        for f in files:
            inputs.append(str(f))
            doc = read_doc(f.name, f.read_bytes())
            saved = {}
            for src, d in by_src.items():
                if not src or f.name in src or src in (doc.get("filename") or "") or any(src == s_.get("label") for s_ in doc.get("sources") or []):
                    saved.update(d)
            cases += doc_cases(f.name, doc, reading, f"test {o['id']}", saved, cust)
        if len(cases) >= limit:
            break
    return cases[:limit], inputs


def from_orders_rows(full: dict) -> list[dict]:
    out = []
    for k, r in enumerate(full["rows"]):
        ln = r.get("ln") or {}
        if not (ln.get("text") or ln.get("desc")):
            continue
        m = r.get("m") or {}
        it = r.get("item") or {}
        opts = [x.get("sku") for x in (m.get("options") or [])]
        ok = r.get("state") == "ok" and not r.get("excluded") and it.get("sku")
        out.append({"src": f"test {full['id']} line {ln.get('n') or k + 1}", "customer": full.get("customer_no") or full.get("customer_name") or "",
                    "text": ln.get("text") or "", "desc": ln.get("desc") or "", "qty": ln.get("qty") or 1, "uom": ln.get("uom") or "",
                    "brand": ln.get("brand") or "", "part": ln.get("part") or "", "from_ocr": bool(ln.get("from_ocr")),
                    "expected": it.get("sku") if ok else None, "how": "", "then_status": m.get("status") or "",
                    "then_top": opts[0] if opts else None, "state": r.get("state") or "", "reading": ""})
    return out


def from_documents(docs: list[tuple[str, bytes]], reading: list) -> list[dict]:
    cases = []
    for name, data in docs:
        cases += doc_cases(name, read_doc(name, data), reading, name)
    return cases


# ------------------------------------------------------------------ matching them
def run(cases: list[dict], branch: str | None = None, level: str | None = None, per_customer: bool = False,
        progress=None, trace: bool = True) -> list[dict]:
    from webbsearch.api.server import bom, customer_key, price_lines
    from webbsearch.retrieval.entitlements import ANONYMOUS
    b = bom()
    out = []
    groups = defaultdict(list)
    for c in cases:
        groups[customer_key(c["customer"]) if per_customer else None].append(c)
    done = 0
    for acct, cs in groups.items():
        for i in range(0, len(cs), 50):
            chunk = cs[i:i + 50]
            lines = [{"text": c["text"], "desc": c["desc"], "qty": c["qty"], "uom": c["uom"], "brand": c["brand"],
                      "part": c["part"], "from_ocr": c["from_ocr"], "edited": not c["from_ocr"]} for c in chunk]
            t0 = time.perf_counter()
            r = b.match_all(lines, ANONYMOUS, branch, k=K, acct=acct, trace=trace)
            price_lines(r, level)
            ms = (time.perf_counter() - t0) * 1000 / max(len(chunk), 1)
            for c, x in zip(chunk, r["lines"]):
                out.append(score(c, x, ms))
            done += len(chunk)
            if progress:
                progress(done, len(cases))
    return out


def score(c: dict, x: dict, ms: float) -> dict:
    opts = x.get("options") or []
    skus = [o.get("sku") for o in opts]
    exp = c.get("expected")
    rank = skus.index(exp) + 1 if exp and exp in skus else None
    top = opts[0] if opts else (x.get("closest") or [{}])[0] if x.get("closest") else {}
    return {**c, "status": x.get("status"), "via": x.get("via") or "", "weak": bool(x.get("weak")), "note": x.get("note") or "",
            "searched_as": x.get("searched_as") or "", "missing": x.get("missing_what") or "", "req_brand": x.get("req_brand") or "",
            "uncarried": bool(x.get("uncarried")), "key": x.get("key") or "", "confirmed": bool(x.get("confirmed")),
            "rank": rank, "hit1": bool(exp) and rank == 1, "hit3": bool(exp) and rank is not None and rank <= 3,
            "top": [{"sku": o.get("sku"), "title": o.get("title"), "brand": o.get("brand"), "cat3": None,
                     "why": "; ".join(w.get("t", "") for w in o.get("why") or []), "stock": (o.get("fulfil") or {}).get("status"),
                     "url": o.get("web_url")} for o in opts[:3]],
            "now_top": top.get("sku"), "changed": bool(c.get("then_top")) and top.get("sku") != c.get("then_top"), "ms": round(ms, 1),
            "steps": steps(x.get("trace") or []), "trace": x.get("trace") or [], "reading": c.get("reading", "")}


def steps(tr: list) -> str:
    """The trace in one line: how the line went through the matcher."""
    out = []
    for t in tr:
        s_ = t.get("step")
        if s_ == "read":
            out.append(f"read: “{t.get('text')}”" + (f" + “{t['desc']}”" if t.get("desc") else "") + f", qty {t.get('qty')}"
                       + (f" {t['uom']}" if t.get("uom") else "") + (f", brand named: {t['brand_named']}" if t.get("brand_named") else ""))
        elif s_ in ("item cross", "brand cross") and t.get("found"):
            out.append(f"{s_}: {t['found']}")
        elif s_ == "exact lookup":
            out.append(f"exact lookup: {t.get('result')}" + (f" ({t['via']})" if t.get("via") else ""))
        elif s_ == "description":
            out.append("searched: “" + (t.get("expanded") or t.get("query") or "") + "”" if t.get("searched") else
                       "not searched (looks like a part number)" if t.get("part_numbers_on_line") else "not searched")
        elif s_ == "search":
            sp = {k: v for k, v in (t.get("spec") or {}).items() if v}
            out.append(f"{t.get('results')} results; line asks for " + ", ".join(f"{k} {v}" for k, v in sp.items()))
        elif s_ == "focused search":
            out.append(f"focused search: “{t.get('query')}” ({t.get('why')})")
        elif s_ == "ranked":
            c = (t.get("candidates") or [{}])[0]
            if c:
                out.append(f"won: {c.get('sku')} (search {c.get('search')} + fit {c.get('fit')} + bonus {c.get('bonus')} = {c.get('total')})")
        elif s_ == "result":
            out.append(f"result: {t.get('status')}" + (f", {t['note']}" if t.get("note") else ""))
    return " → ".join(out)


# ------------------------------------------------------------------ the numbers
def summarize(rows: list[dict], engine) -> dict:
    n = len(rows)
    st = Counter(r["status"] for r in rows)
    withexp = [r for r in rows if r["expected"]]
    cat = engine.catalog

    def cat3(sku):
        j = cat.sku_to_idx.get(sku)
        return str(cat.cat3[j]) if j is not None else "(not on the list)"
    by = defaultdict(lambda: [0, 0, 0])
    for r in withexp:
        g = by[cat3(r["expected"])]
        g[0] += 1
        g[1] += r["hit1"]
        g[2] += r["hit3"]
    shape = defaultdict(lambda: [0, 0])
    for r in withexp:
        k = "part number" if re.search(r"\b(?=[A-Z0-9-]*\d)(?=[A-Z0-9-]*[A-Z])[A-Z0-9-]{4,}\b", r["text"]) else \
            "handwriting" if r["from_ocr"] else "description"
        shape[k][0] += 1
        shape[k][1] += r["hit1"]
    mrr = sum(1 / r["rank"] for r in withexp if r["rank"]) / len(withexp) if withexp else None
    return {"lines": n, "status": dict(st), "matched_rate": round(1 - st.get("none", 0) / n, 4) if n else None,
            "weak_rate": round(sum(r["weak"] for r in rows) / n, 4) if n else None,
            "with_expected": len(withexp),
            "top1": round(sum(r["hit1"] for r in withexp) / len(withexp), 4) if withexp else None,
            "top3": round(sum(r["hit3"] for r in withexp) / len(withexp), 4) if withexp else None,
            "in_options": round(sum(bool(r["rank"]) for r in withexp) / len(withexp), 4) if withexp else None,
            "mrr": round(mrr, 4) if mrr is not None else None,
            "changed_since_shown": sum(r["changed"] for r in rows),
            "how_staff_chose": dict(Counter(r["how"] for r in rows if r["how"])),
            "by_category": sorted(([k, v[0], round(v[1] / v[0], 3), round(v[2] / v[0], 3)] for k, v in by.items()), key=lambda x: -x[1])[:40],
            "by_line_shape": {k: [v[0], round(v[1] / v[0], 3)] for k, v in shape.items()},
            "ms_per_line": round(sum(r["ms"] for r in rows) / n, 1) if n else None,
            "not_found": Counter((r["missing"] or r["text"]).lower() for r in rows if r["status"] == "none").most_common(60),
            "brands_not_carried": Counter(r["req_brand"] for r in rows if r["uncarried"] and r["req_brand"]).most_common(40)}


# ------------------------------------------------------------------ the workbook
ABOUT = [
    ("Summary", "The numbers for this run. top1 / top3: of lines with an expected item, the share where the matcher ranked it first / in the first three. in_options: anywhere in the first 8. mrr: mean reciprocal rank (1 = always first). matched_rate: lines with any match."),
    ("Trace", "Every candidate the matcher ranked for each line (first 15): search position, whether it is the product the line names, the right size, the spec, and the points (search + fit + bonus, with the bonus parts). Rank order: the product first, then its size, then nothing missing, then total points."),
    ("Reading", "Documents only (tests, documents): every line read from each document, lines set aside and why, and for test cases whether the text and quantity were read as staff corrected them."),
    ("Lines", "Every line. expected: the item it should match (a confirmed order line, a staff choice in the log, or the test file's answer). rank: where the matcher put it now (blank: not in the first 8). then_top: what was shown at the time. top1..3: today's first three, with the reasons shown to staff."),
    ("Misses", "Lines with an expected item the matcher didn't rank first: the place to look for fixes (a dictionary entry, a cross, a brand order, a weight)."),
    ("Not found", "Lines with no match, grouped by what was asked for."),
    ("Brands", "Brands named on lines that aren't on the item list: candidates for brand crosses."),
    ("Changed", "Lines whose top match today differs from what staff saw then (settings or the item list changed)."),
    ("Settings", "The bulk weights, crosses and brand orders in force when the report ran."),
]


def workbook(meta: dict, summary: dict, rows: list[dict], engine, reading: list | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    cat = engine.catalog

    def title(sku):
        j = cat.sku_to_idx.get(sku) if sku else None
        return str(cat.title[j]) if j is not None else ""
    wb = Workbook()
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8ECEF")

    def sheet(name, head, data, widths):
        ws = wb.create_sheet(name)
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = bold, fill
        for row in data:
            ws.append([("'" + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v for v in row])
        for k, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(k)].width = w
        ws.freeze_panes = "A2"
        if data:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(data) + 1}"
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
        return ws

    s = wb.active
    s.title = "Summary"
    for k, v in [("Report", meta["id"]), ("Source", meta["source_label"]), ("Ran", meta["at"]), ("Items on the list", f"{cat.n:,}"),
                 ("Branch", meta.get("branch") or "default"), ("Per-customer learning", "on" if meta.get("per_customer") else "off"),
                 ("Settings version", meta.get("settings_version"))]:
        s.append([k, v])
    s.append([])
    prev = meta.get("previous") or {}
    s.append(["Measure", "This run", "Previous run", "Change"])
    for k in ("lines", "with_expected", "top1", "top3", "in_options", "mrr", "matched_rate", "weak_rate", "changed_since_shown", "ms_per_line"):
        v, p = summary.get(k), prev.get(k)
        s.append([k, v, p, round(v - p, 4) if isinstance(v, (int, float)) and isinstance(p, (int, float)) else None])
    for k, v in (summary.get("reading") or {}).items():
        s.append([f"reading: {k}", v])
    s.append([])
    s.append(["Status", "Lines"])
    for k, v in sorted(summary["status"].items(), key=lambda x: -x[1]):
        s.append([k, v])
    s.append([])
    s.append(["How staff chose (orders and log)", "Lines"])
    for k, v in summary["how_staff_chose"].items():
        s.append([k, v])
    s.append([])
    s.append(["Line shape", "Lines with expected", "Top-1"])
    for k, v in summary["by_line_shape"].items():
        s.append([k, v[0], v[1]])
    s.append([])
    s.append(["Category of the expected item", "Lines", "Top-1", "Top-3"])
    for row in summary["by_category"]:
        s.append(row)
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 36, 22, 14
    for r in s.iter_rows():
        if r[0].value in ("Measure", "Status", "How staff chose (orders and log)", "Line shape", "Category of the expected item"):
            for c in r:
                c.font, c.fill = bold, fill

    def tops(r, i):
        t = r["top"][i] if len(r["top"]) > i else {}
        return [t.get("sku"), t.get("title"), t.get("why")]
    sheet("Lines", ["Source", "Customer", "Requested", "Qty", "UOM", "Status", "Via", "Weak", "Note", "Searched as", "Expected item",
                    "Expected description", "Rank of expected", "How staff chose", "Shown then", "Changed",
                    "Top 1 item", "Top 1 description", "Top 1 reasons", "Top 2 item", "Top 2 description", "Top 2 reasons",
                    "Top 3 item", "Top 3 description", "Top 3 reasons", "Line key", "ms", "Reading check", "Steps"],
          [[r["src"], r["customer"], r["text"] + (f" | {r['desc']}" if r["desc"] else ""), r["qty"], r["uom"], r["status"], r["via"],
            "yes" if r["weak"] else "", r["note"], r["searched_as"], r["expected"], title(r["expected"]), r["rank"], r["how"],
            r["then_top"], "yes" if r["changed"] else "", *tops(r, 0), *tops(r, 1), *tops(r, 2), r["key"], r["ms"],
            r.get("reading") or "", r["steps"]] for r in rows],
          [22, 14, 36, 6, 6, 8, 16, 6, 30, 24, 12, 40, 8, 18, 12, 8, 12, 44, 30, 12, 36, 24, 12, 36, 24, 24, 6, 30, 90])
    miss = [r for r in rows if r["expected"] and not r["hit1"]]
    sheet("Misses", ["Requested", "Qty", "Expected item", "Expected description", "Rank of expected", "Got (top 1)", "Got description",
                     "Got reasons", "Status", "Note", "Searched as", "Source"],
          [[r["text"], r["qty"], r["expected"], title(r["expected"]), r["rank"] or "not in top 8", *tops(r, 0), r["status"], r["note"],
            r["searched_as"], r["src"]] for r in miss],
          [36, 6, 12, 44, 10, 12, 44, 30, 8, 30, 24, 22])
    trace_rows = []
    for r in rows:
        for t in r["trace"]:
            if t.get("step") == "ranked":
                exp = r["expected"]
                for c in t.get("candidates") or []:
                    trace_rows.append([r["text"], c.get("rank"), c.get("sku"), c.get("title"), c.get("brand"), "yes" if c.get("sku") == exp else "",
                                       c.get("search_pos"), "yes" if c.get("product") else "no", "yes" if c.get("size") else "no",
                                       "yes" if c.get("spec_ok") else "no", c.get("misses"), c.get("search"), c.get("fit"), c.get("bonus"),
                                       "; ".join(f"{k} {v}" for k, v in (c.get("bonus_parts") or {}).items()), c.get("total")])
    sheet("Trace", ["Requested", "Rank", "Item", "Description", "Brand", "Expected", "Search position", "Is the product",
                    "Right size", "Spec ok", "Misses", "Search pts", "Fit pts", "Bonus pts", "Bonus parts", "Total"],
          trace_rows, [32, 6, 12, 48, 16, 8, 8, 8, 8, 8, 20, 8, 8, 8, 30, 8])
    if reading:
        sheet("Reading", ["Document", "Source", "Line", "Raw", "Read as text", "Qty", "UOM", "Read as", "Why / confidence", "Check"],
              [[x["doc"], x["source"], x["n"], x["raw"], x["text"], x["qty"], x["uom"], x["read_as"], x["why"], x["check"]] for x in reading],
              [24, 24, 6, 44, 36, 6, 6, 18, 24, 40])
    sheet("Not found", ["Asked for", "Lines"], summary["not_found"], [50, 8])
    sheet("Brands", ["Brand on the line", "Lines"], summary["brands_not_carried"], [30, 8])
    ch = [r for r in rows if r["changed"]]
    sheet("Changed", ["Requested", "Shown then", "Then description", "Now", "Now description", "Expected", "Source"],
          [[r["text"], r["then_top"], title(r["then_top"]), r["now_top"], title(r["now_top"]), r["expected"], r["src"]] for r in ch],
          [36, 12, 40, 12, 40, 12, 22])
    ms = engine.merch.get()
    st = wb.create_sheet("Settings")
    st.append(["Bulk weight", "Value"])
    for k, v in (ms.get("bulk") or {}).items():
        st.append([k, v])
    st.append([])
    st.append(["Cross", "From", "To", "Where", "On"])
    for c in ms.get("crosses") or []:
        st.append([c.get("kind"), c.get("from_brand") or c.get("from_part") or c.get("from_text"), c.get("to_brand") or c.get("to_sku"),
                   c.get("scope"), c.get("enabled")])
    st.append([])
    st.append(["Brand order", "Brands", "Where", "On"])
    for b in ms.get("brand_prefs") or []:
        st.append([b.get("family"), ", ".join(b.get("brands") or []), b.get("scope"), b.get("enabled")])
    st.column_dimensions["A"].width, st.column_dimensions["B"].width, st.column_dimensions["C"].width = 26, 40, 24
    ab = wb.create_sheet("About")
    ab.append(["Sheet", "What it holds"])
    for row in ABOUT:
        ab.append(list(row))
    ab.append([])
    ab.append(["Summary data", json.dumps({"meta": meta, "summary": summary}, default=str)[:30000]])
    ab.column_dimensions["A"].width, ab.column_dimensions["B"].width = 14, 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ one report, start to end
def make(source: str, days: int = 90, file_name: str | None = None, file_data: bytes | None = None, branch: str | None = None,
         level: str | None = None, per_customer: bool = False, progress=None, documents: list | None = None) -> dict:
    from webbsearch.api.server import engine
    e = engine()
    reading: list = []
    inputs: list = []
    if source == "tests":
        cases, inputs = from_tests(reading)
        label = "test cases"
    elif source == "documents":
        if not documents:
            raise ValueError("no documents given")
        cases, label = from_documents(documents, reading), f"{len(documents)} customer document{'s' if len(documents) > 1 else ''}"
    elif source == "orders":
        cases, label = from_orders(days), f"saved orders, last {days} days"
    elif source == "log":
        cases, label = from_log(days), f"lines matched in the app, last {days} days"
    elif source == "file":
        if not file_data:
            raise ValueError("a test file is needed")
        cases, label = from_file(file_name or "test.csv", file_data), f"test file {file_name}"
    else:
        raise ValueError(f"source is one of {', '.join(SOURCES)}")
    if not cases:
        raise ValueError(f"no lines in {label}")
    t0 = time.time()
    rows = run(cases, branch, level, per_customer, progress)
    summary = summarize(rows, e)
    if reading:
        chk = [r for r in reading if r["check"]]
        summary["reading"] = {"lines_read": sum(1 for r in reading if r["read_as"].startswith("item")),
                              "set_aside": sum(1 for r in reading if r["read_as"] == "set aside"),
                              "not_read": sum(1 for r in reading if r["read_as"] == "NOT READ"),
                              "checked": len(chk), "read_right": sum(1 for r in chk if r["check"] == "ok")}
    REPORTS.mkdir(parents=True, exist_ok=True)
    rid = time.strftime("R%y%m%d-%H%M-") + secrets.token_hex(2).upper()
    prev = previous(source, file_name)
    meta = {"id": rid, "source": source, "source_label": label, "file": file_name, "days": days, "branch": branch,
            "level": level, "per_customer": per_customer, "at": time.strftime("%Y-%m-%d %H:%M"), "secs": round(time.time() - t0, 1),
            "settings_version": e.merch.get().get("version"), "items": e.catalog.n,
            "previous": {k: prev["summary"].get(k) for k in ("lines", "with_expected", "top1", "top3", "in_options", "mrr",
                                                             "matched_rate", "weak_rate", "changed_since_shown", "ms_per_line")}
            | {"id": prev["meta"]["id"]} if prev else None}
    meta["inputs"] = inputs
    if documents:                                  # kept with the report, for the bundle
        d = REPORTS / f"{rid}-inputs"
        d.mkdir(exist_ok=True)
        for name, data in documents:
            (d / re.sub(r"[^\w.() -]+", "_", Path(name).name)[:120]).write_bytes(data)
    elif file_data and file_name:
        d = REPORTS / f"{rid}-inputs"
        d.mkdir(exist_ok=True)
        (d / re.sub(r"[^\w.() -]+", "_", Path(file_name).name)[:120]).write_bytes(file_data)
    (REPORTS / f"{rid}.xlsx").write_bytes(workbook(meta, summary, rows, e, reading))
    (REPORTS / f"{rid}.json").write_text(json.dumps({"meta": meta, "summary": summary}, default=str))
    return {"meta": meta, "summary": {k: v for k, v in summary.items() if k not in ("not_found", "brands_not_carried", "by_category")}}


def previous(source: str, file_name: str | None) -> dict | None:
    best = None
    for f in sorted(REPORTS.glob("R*.json")) if REPORTS.exists() else []:
        try:
            d = json.loads(f.read_text())
        except ValueError:
            continue
        if d["meta"]["source"] == source and (source != "file" or d["meta"].get("file") == file_name):
            best = d
    return best


def listing(limit: int = 50) -> list[dict]:
    out = []
    for f in sorted(REPORTS.glob("R*.json"), reverse=True)[:limit] if REPORTS.exists() else []:
        try:
            d = json.loads(f.read_text())
        except ValueError:
            continue
        s = d["summary"]
        out.append({"id": d["meta"]["id"], "at": d["meta"]["at"], "source": d["meta"]["source_label"], "lines": s["lines"],
                    "with_expected": s["with_expected"], "top1": s["top1"], "top3": s["top3"], "matched_rate": s["matched_rate"],
                    "secs": d["meta"].get("secs")})
    return out


def mac_libomp(module: str) -> None:
    """macOS: LightGBM needs the OpenMP library scikit-learn ships; start.sh points the loader at it, and a command
    run by hand gets the same by starting itself again with the path set (once)."""
    import sys
    if sys.platform != "darwin" or os.environ.get("FLEXBULK_OMP"):
        return
    import sklearn
    lib = os.path.join(os.path.dirname(sklearn.__file__), ".dylibs")
    env = dict(os.environ, FLEXBULK_OMP="1",
               DYLD_FALLBACK_LIBRARY_PATH=lib + (":" + os.environ["DYLD_FALLBACK_LIBRARY_PATH"] if os.environ.get("DYLD_FALLBACK_LIBRARY_PATH") else ""))
    os.execve(sys.executable, [sys.executable, "-m", module, *sys.argv[1:]], env)


def main(argv=None):
    mac_libomp("flexbulk.report")
    ap = argparse.ArgumentParser(description="Write a match report (Excel) from saved orders, the app's log or a test file")
    ap.add_argument("--source", choices=SOURCES, default="orders")
    ap.add_argument("--days", type=int, default=90)
    ap.add_argument("--file")
    ap.add_argument("--branch")
    ap.add_argument("--level")
    ap.add_argument("--per-customer", action="store_true")
    a = ap.parse_args(argv)
    data = Path(a.file).read_bytes() if a.file else None
    r = make(a.source, a.days, Path(a.file).name if a.file else None, data, a.branch, a.level, a.per_customer,
             progress=lambda d, n: print(f"  {d:,} of {n:,} lines", flush=True) if d % 500 == 0 or d == n else None)
    print(json.dumps(r["summary"], indent=1))
    print(f"report: {REPORTS / (r['meta']['id'] + '.xlsx')}")


if __name__ == "__main__":
    main()


# ------------------------------------------------------------------ the bundle: the report and what it takes to fix things
CODE = ["webbsearch/retrieval/bom.py", "webbsearch/retrieval/resolve.py", "webbsearch/retrieval/linefit.py",
        "webbsearch/retrieval/intake.py", "webbsearch/retrieval/shorthand.py", "webbsearch/retrieval/trade_terms.py",
        "webbsearch/retrieval/availability.py", "webbsearch/merch/bulkprefs.py", "webbsearch/merch/rules.py",
        "webbsearch/merch/store.py", "webbsearch/api/server.py", "webbsearch/api/static/bulk.js",
        "webbsearch/api/static/bulkapp.js", "webbsearch/api/static/index.html", "flexbulk/loader.py", "flexbulk/report.py",
        "flexbulk/prices.py", "config/bulk_sources.yaml"]
BUNDLE_README = """# Flex Bulk match report {rid}

Everything needed to work on the matching from this report, without the repository:

- report.xlsx              the report (start with Summary, then Misses; Trace shows how every candidate scored)
- settings.json            the dashboard settings in force (weights, crosses, brand orders, learned-cross switches)
- code/                    the matcher and the staff page as they are (Flex Bulk {version})
    webbsearch/retrieval/bom.py       the matcher: read -> crosses -> exact lookup -> search -> rank (recommend())
    webbsearch/retrieval/resolve.py   exact item #, mfr part #, UPC lookups
    webbsearch/retrieval/linefit.py   does a candidate fit the line: product, size, material, joint
    webbsearch/retrieval/intake.py    reading emails, Word, Excel, PDF, photos into lines
    webbsearch/retrieval/shorthand.py trade shorthand (cxc, bv, 45/5...) and OCR repairs
    webbsearch/merch/                 crosses, brand orders, learned choices, bulk weights
    webbsearch/api/static/            the staff page (bulk.js: the lines table; bulkapp.js: the page around it)
    flexbulk/                         loading the item list, reports, prices
    config/bulk_sources.yaml          how the item file's columns are read
- inputs/                  the documents and test file the lines came from{inputs_note}

How a line is ranked (bom.py, recommend): candidates from the search are sorted by
(not the product the line names, not its size, something it asked for missing, -points), where
points = search position (10 for first, falling) + fit (linefit) + bonus (on the shelf, Webb's brand order, learned picks).
The item list has {items:,} items; the search itself (webbsearch/retrieval/results.py and the indexes) is not included.
"""


def bundle(rid: str, inputs: bool = True) -> bytes:
    import zipfile
    from webbsearch import __version__
    from webbsearch.config import ROOT
    from webbsearch.api.server import engine
    meta = json.loads((REPORTS / f"{rid}.json").read_text())["meta"]
    buf = io.BytesIO()
    files = []
    if inputs:
        d = REPORTS / f"{rid}-inputs"
        files = [(f"inputs/{f.name}", f) for f in sorted(d.iterdir())] if d.exists() else []
        files += [(f"inputs/{Path(x).parent.name}/{Path(x).name}", Path(x)) for x in meta.get("inputs") or [] if Path(x).exists()]
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.md", BUNDLE_README.format(rid=rid, version=__version__, items=meta.get("items") or 0,
                                                     inputs_note="" if inputs else " (left out of this download)"))
        z.write(REPORTS / f"{rid}.xlsx", "report.xlsx")
        z.writestr("settings.json", json.dumps(engine().merch.get(), indent=1, default=str))
        for rel in CODE:
            f = ROOT / rel
            if f.exists():
                z.write(f, f"code/{rel}")
        for arc, f in files:
            z.write(f, arc)
    return buf.getvalue()
