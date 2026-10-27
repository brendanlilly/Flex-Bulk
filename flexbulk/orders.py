"""Saved orders: who it's for (customer, job, PO), who prepared it, and the matched lines as staff left them.

One SQLite file (data/bulk/orders.sqlite). An order is saved whole each time (the page sends every line with its
match and the choice made), so reopening it shows exactly what was there: the options, the pick, what was confirmed.
Exports (Excel, CSV, and a two-column item/quantity file for order entry) are built from the saved copy.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import secrets
import sqlite3
import threading
import time
from pathlib import Path

from webbsearch.config import settings

HEAD = ("customer_no", "customer_name", "job", "po", "contact", "branch", "branch_name", "level", "staff", "notes")
MAX_LINES = 2000
STATUS_LABEL = {"ok": "Confirmed", "confirm": "To confirm", "none": "Not found", "out": "Left out", "skip": "Skipped",
                "new": "", "busy": ""}


def _docs(body: dict) -> str:
    """The documents the lines came from (names, what they said about the order); dropped if oversized."""
    d = json.dumps(body.get("docs") or [])
    return d if len(d) <= 400_000 else "[]"


def _clean(v, n: int = 200) -> str:
    return re.sub(r"[\x00-\x1f]", " ", str(v or "")).strip()[:n]


class Orders:
    def __init__(self, path: Path | None = None):
        self.path = Path(path or os.environ.get("FLEXBULK_ORDERS_DB") or settings.data_dir / "bulk" / "orders.sqlite")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.Lock()
        self.db = sqlite3.connect(self.path, check_same_thread=False)
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.execute("""CREATE TABLE IF NOT EXISTS orders (
            id TEXT PRIMARY KEY, created REAL, updated REAL, status TEXT, customer_no TEXT, customer_name TEXT,
            job TEXT, po TEXT, contact TEXT, branch TEXT, branch_name TEXT, level TEXT, staff TEXT, notes TEXT,
            n_lines INTEGER, n_confirmed INTEGER, total REAL, rows TEXT, docs TEXT)""")
        self.db.commit()

    # ------------------------------------------------------------------ write
    def save(self, body: dict) -> dict:
        """Create (no id) or replace an order. Returns the saved order's summary."""
        rows = [self._row(r) for r in (body.get("rows") or [])[:MAX_LINES]]
        rows = [r for r in rows if r["text"] or r["item"]]
        head = {k: _clean(body.get(k), 2000 if k == "notes" else 200) for k in HEAD}
        conf = [r for r in rows if r["state"] == "ok"]
        total = round(sum((r["item"] or {}).get("ext") or 0 for r in conf), 2)
        status = body.get("status") if body.get("status") in ("draft", "done") else "draft"
        now = time.time()
        oid = body.get("id")
        with self.lock:
            if oid and self.db.execute("SELECT 1 FROM orders WHERE id=?", (oid,)).fetchone():
                self.db.execute(f"UPDATE orders SET updated=?, status=?, {', '.join(f'{k}=?' for k in HEAD)}, n_lines=?, "
                                "n_confirmed=?, total=?, rows=?, docs=? WHERE id=?",
                                (now, status, *[head[k] for k in HEAD], len(rows), len(conf), total, json.dumps(rows),
                                 _docs(body), oid))
            else:
                oid = time.strftime("B%y%m%d-") + secrets.token_hex(2).upper()
                self.db.execute(f"INSERT INTO orders (id, created, updated, status, {', '.join(HEAD)}, n_lines, n_confirmed, "
                                f"total, rows, docs) VALUES (?,?,?,?,{','.join('?' * len(HEAD))},?,?,?,?,?)",
                                (oid, now, now, status, *[head[k] for k in HEAD], len(rows), len(conf), total,
                                 json.dumps(rows), _docs(body)))
            self.db.commit()
        return self.summary(oid)

    @staticmethod
    def _row(r: dict) -> dict:
        """One line as saved: what was asked, the state, the item chosen (a snapshot), and the match (to reopen)."""
        ln = r.get("ln") or {}
        it = r.get("item") or None
        m = r.get("m") or None
        if m:                         # handwriting crops and other bulky bits stay out of the database
            m = {k: v for k, v in m.items() if k not in ("crop", "ocr_raw")}
        qty = max(1, int(float(ln.get("qty") or 1)))
        item = None
        if it:
            units = int(it.get("units") or qty)
            price = it.get("price")
            item = {"sku": _clean(it.get("sku"), 60), "title": _clean(it.get("title"), 300), "brand": _clean(it.get("brand"), 80),
                    "mfr_part": _clean(it.get("mfr_part"), 80), "uom": _clean(it.get("uom") or "EA", 8), "units": units,
                    "conv": _clean(it.get("conv"), 80), "price": price if isinstance(price, (int, float)) else None,
                    "ext": round(price * units, 2) if isinstance(price, (int, float)) else None,
                    "stock": _clean((it.get("fulfil") or {}).get("text"), 300),
                    "stock_status": _clean((it.get("fulfil") or {}).get("status"), 20),
                    "why": "; ".join(_clean(w.get("t"), 120) for w in (it.get("why") or []) if isinstance(w, dict))}
        return {"n": ln.get("n"), "text": _clean(ln.get("text"), 300), "desc": _clean(ln.get("desc"), 300),
                "qty": qty, "uom": _clean(ln.get("uom"), 8), "src": _clean(ln.get("src_label"), 120),
                "state": r.get("state") if r.get("state") in STATUS_LABEL else "new", "item": item,
                "note": _clean((m or {}).get("note"), 300), "ln": {k: v for k, v in ln.items() if k != "crop"},
                "m": m, "ch": r.get("ch") or {"t": "opt", "i": 0}, "confirmed": bool(r.get("confirmed")),
                "excluded": bool(r.get("excluded"))}

    def delete(self, oid: str) -> bool:
        with self.lock:
            n = self.db.execute("DELETE FROM orders WHERE id=?", (oid,)).rowcount
            self.db.commit()
        return bool(n)

    # ------------------------------------------------------------------ read
    COLS = ("id", "created", "updated", "status", *HEAD, "n_lines", "n_confirmed", "total")

    def summary(self, oid: str) -> dict | None:
        r = self.db.execute(f"SELECT {', '.join(self.COLS)} FROM orders WHERE id=?", (oid,)).fetchone()
        return dict(zip(self.COLS, r)) if r else None

    def get(self, oid: str) -> dict | None:
        r = self.db.execute(f"SELECT {', '.join(self.COLS)}, rows, docs FROM orders WHERE id=?", (oid,)).fetchone()
        if not r:
            return None
        d = dict(zip(self.COLS, r[:-2]))
        d["rows"] = json.loads(r[-2] or "[]")
        d["docs"] = json.loads(r[-1] or "[]")
        return d

    def list(self, q: str = "", limit: int = 100, customer: str = "") -> list[dict]:
        sql = f"SELECT {', '.join(self.COLS)} FROM orders"
        args: list = []
        where = []
        if q:
            like = f"%{q.strip()}%"
            where.append("(customer_no LIKE ? OR customer_name LIKE ? OR job LIKE ? OR po LIKE ? OR id LIKE ? OR staff LIKE ?)")
            args += [like] * 6
        if customer:
            where.append("customer_no = ?")
            args.append(customer)
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY updated DESC LIMIT ?"
        args.append(max(1, min(int(limit), 500)))
        return [dict(zip(self.COLS, r)) for r in self.db.execute(sql, args).fetchall()]

    def customers(self, q: str = "", limit: int = 12) -> list[dict]:
        """Customers staff have entered before (for the customer box's suggestions)."""
        like = f"%{q.strip()}%"
        rows = self.db.execute("SELECT customer_no, customer_name, MAX(updated), COUNT(*) FROM orders "
                               "WHERE (customer_no LIKE ? OR customer_name LIKE ?) AND (customer_no<>'' OR customer_name<>'') "
                               "GROUP BY customer_no, customer_name ORDER BY MAX(updated) DESC LIMIT ?",
                               (like, like, limit)).fetchall()
        return [{"customer_no": a, "customer_name": b, "orders": n} for a, b, _, n in rows]


# ------------------------------------------------------------------ exports
def _safe(v):
    """A cell that would run as a spreadsheet formula gets a leading apostrophe (customer text is untrusted)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") and not re.fullmatch(r"[+-]?\d+(\.\d+)?", v):
        return "'" + v
    return v


LINE_HEAD = ["Line", "Requested", "Description", "Qty", "UOM", "Status", "Item #", "Manufacturer", "Mfr part #",
             "Item description", "Order qty", "Unit price", "Extended", "Stock", "Why this item", "Note", "Source"]


def line_rows(o: dict) -> list[list]:
    out = []
    for k, r in enumerate(o["rows"]):
        it = r.get("item") or {}
        st = "Left out" if r.get("excluded") else STATUS_LABEL.get(r.get("state"), "")
        out.append([r.get("n") or k + 1, r.get("text"), r.get("desc") if r.get("desc") != r.get("text") else "", r.get("qty"),
                    r.get("uom") or "", st, it.get("sku", ""), it.get("brand", ""), it.get("mfr_part", ""), it.get("title", ""),
                    it.get("units", "") if it else "", it.get("price"), it.get("ext"), it.get("stock", ""), it.get("why", ""),
                    r.get("note", ""), r.get("src", "")])
    return out


def head_rows(o: dict) -> list[tuple]:
    when = time.strftime("%Y-%m-%d %H:%M", time.localtime(o["updated"]))
    return [("Order", o["id"]), ("Customer", " · ".join(x for x in (o["customer_no"], o["customer_name"]) if x)),
            ("Job", o["job"]), ("PO", o["po"]), ("Contact", o["contact"]), ("Pickup branch", o["branch_name"] or o["branch"]),
            ("Price level", o["level"] or "List"), ("Prepared by", o["staff"]), ("Saved", when), ("Notes", o["notes"])]


def to_csv(o: dict) -> bytes:
    f = io.StringIO()
    w = csv.writer(f)
    for k, v in head_rows(o):
        if v:
            w.writerow([k, _safe(v)])
    w.writerow([])
    w.writerow(LINE_HEAD)
    for r in line_rows(o):
        w.writerow([_safe(v) if v is not None else "" for v in r])
    return f.getvalue().encode("utf-8-sig")


def to_entry_csv(o: dict) -> bytes:
    """Confirmed lines only, item # and order quantity: the file order entry imports."""
    f = io.StringIO()
    w = csv.writer(f)
    w.writerow(["Item #", "Qty"])
    for r in o["rows"]:
        it = r.get("item")
        if it and r.get("state") == "ok" and not r.get("excluded"):
            w.writerow([_safe(it["sku"]), it.get("units") or r.get("qty")])
    return f.getvalue().encode("utf-8-sig")


def to_xlsx(o: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Order"
    bold = Font(bold=True)
    r = 1
    for k, v in head_rows(o):
        if v:
            ws.cell(r, 1, k).font = bold
            ws.cell(r, 2, _safe(v))
            r += 1
    r += 1
    head_row = r
    fill = PatternFill("solid", fgColor="E8ECEF")
    for c, h in enumerate(LINE_HEAD, 1):
        cell = ws.cell(r, c, h)
        cell.font = bold
        cell.fill = fill
    for line in line_rows(o):
        r += 1
        for c, v in enumerate(line, 1):
            ws.cell(r, c, _safe(v) if v is not None else None)
        for c in (12, 13):
            ws.cell(r, c).number_format = '"$"#,##0.00'
    conf = [x for x in o["rows"] if x.get("state") == "ok" and not x.get("excluded") and x.get("item")]
    r += 2
    ws.cell(r, 10, "Confirmed lines").font = bold
    ws.cell(r, 11, len(conf))
    ws.cell(r, 12, "Total").font = bold
    ws.cell(r, 13, round(sum((x["item"].get("ext") or 0) for x in conf), 2)).number_format = '"$"#,##0.00'
    widths = [6, 34, 24, 7, 6, 11, 12, 16, 16, 44, 9, 11, 12, 34, 30, 30, 18]
    for c, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = wdt
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(LINE_HEAD))}{max(head_row, r - 2)}"
    for row in ws.iter_rows(min_row=head_row + 1, max_row=r):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 10, 14, 15, 16))
    e = wb.create_sheet("Order entry")
    e.append(["Item #", "Qty"])
    for x in conf:
        e.append([_safe(x["item"]["sku"]), x["item"].get("units") or x.get("qty")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_store: Orders | None = None


def store() -> Orders:
    global _store
    if _store is None:
        _store = Orders()
    return _store
